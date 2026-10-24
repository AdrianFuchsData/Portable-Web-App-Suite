"""Round trips and conformance with other software: each app opens a file made by other software (tests/helpers/corpus.py),
saves it in every format it writes, and the result is checked by programs that know the format – LibreOffice, 7-Zip,
unzip, ffprobe, pdfinfo/pdftotext, xmllint, SQLite, and the Python libraries python-docx, openpyxl, python-pptx, icalendar,
pykeepass, nbformat, Pillow, psd-tools and pypdf. The content must survive: texts, values, formulas, entries, events.
Formats whose checking program is missing are skipped."""
import io
import json
import re
import shutil
import sqlite3
import tarfile
import subprocess
import wave
import zipfile
from pathlib import Path

import pytest
from playwright.sync_api import expect

from helpers import corpus
from helpers.hostile import drop_file

NO_PICKERS = "window.showOpenFilePicker = undefined; window.showSaveFilePicker = undefined;"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


@pytest.fixture(autouse=True)
def desktop_only(request):
    if request.config.getoption("--mobile"):
        pytest.skip("file formats do not depend on the display")


def tool(name):
    exe = shutil.which(name)
    if not exe:
        pytest.skip(f"{name} is not installed")
    return exe


def run(*args, timeout=180):
    return subprocess.run([str(a) for a in args], capture_output=True, text=True, timeout=timeout)


def soffice_convert(path, target, tmp_path):
    """LibreOffice converts the file (it must be able to read it); the converted file."""
    if not SOFFICE:
        pytest.skip("LibreOffice is not installed")
    out = tmp_path / f"lo-{path.stem}-{target.split(':')[0]}"
    out.mkdir(exist_ok=True)
    run(SOFFICE, f"-env:UserInstallation=file://{tmp_path}/lo-profile", "--headless", "--convert-to", target, "--outdir", out, path)
    files = list(out.iterdir())
    assert files, f"LibreOffice could not read {path.name}"
    return files[0]


def words(text):
    return re.sub(r"\s+", " ", text).strip()


def well_formed_xml_parts(path):
    """Every XML part of a ZIP based file is well-formed (xmllint)."""
    exe = tool("xmllint")
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                r = subprocess.run([exe, "--noout", "-"], input=z.read(name), capture_output=True, timeout=60)
                assert r.returncode == 0, f"{path.name}:{name}: {r.stderr.decode()[:300]}"


def download(page, action, tmp_path, prefix="saved-"):
    with page.expect_download() as d:
        action()
    path = tmp_path / (prefix + d.value.suggested_filename)
    d.value.save_as(path)
    return path


def save_as(page, fmt, tmp_path, ok="[data-testid=save-as-ok]", select="[data-testid=save-as-format]", button="[data-testid=save-as]"):
    page.click(button)
    page.select_option(select, fmt)
    return download(page, lambda: page.click(ok), tmp_path)


def open_dialog(page, path, button="[data-testid=open]"):
    with page.expect_file_chooser() as chooser:
        page.click(button)
    chooser.value.set_files(str(path))


def seed(tmp_path, data, name):
    path = tmp_path / name
    path.write_bytes(data)
    return path


# ---- text and binary files: saved unchanged, byte for byte ---------------------------------------------------------------

TEXTS = {
    "utf16.txt": "﻿Zeile 1 – ä\r\nZeile 2\r\n".encode("utf-16-le").replace(b"\xff\xfe\xff\xfe", b"\xff\xfe"),
    "latin1.txt": "Grüße aus Österreich\nnoch eine Zeile".encode("cp1252"),
    "bom.csv": "﻿a;b\r\n1;2\r\n".encode("utf-8"),
    "mixed.md": b"# Title\r\n\r\ntext\ttab\r\n\r\n",
    "noeol.py": b"print('x')",
}


def test_text_editor_saves_files_unchanged(app_page, tmp_path):
    page = app_page("texteditor", init_script=NO_PICKERS)
    for name, data in TEXTS.items():
        page.evaluate("([n, b]) => TE.app.openPicked([{ file: new File([new Uint8Array(b)], n), handle: null }])", [name, list(data)])
        page.wait_for_function("(n) => TE.app.docs[TE.app.active].name === n", arg=name)
        saved = download(page, lambda: page.keyboard.press("Control+s"), tmp_path)
        assert saved.read_bytes() == data, name
    assert page.errors == []


def test_hex_editor_saves_files_unchanged(app_page, tmp_path):
    page = app_page("hexeditor", init_script=NO_PICKERS)
    data = bytes(range(256)) * 300 + b"\0" * 7
    page.evaluate("([n, b]) => HE.app.load(n, new Uint8Array(b))", ["blob.bin", list(data)])
    saved = download(page, lambda: page.click("[data-testid=save]"), tmp_path)
    assert saved.read_bytes() == data
    assert page.errors == []


# ---- office --------------------------------------------------------------------------------------------------------------

DOC_TEXTS = ["Seed document", "Plain, bold and italic.", "Item one", "Item two", "Second page"]


def test_document_round_trip(app_page, tmp_path):
    from docx import Document
    page = app_page("document", init_script=NO_PICKERS)
    open_dialog(page, seed(tmp_path, corpus.docx(), "seed.docx"))
    expect(page.locator("[data-testid=file-name]")).to_contain_text("seed.docx", timeout=30000)
    page.wait_for_timeout(1000)
    for fmt in ("docx", "odt"):
        saved = save_as(page, fmt, tmp_path)
        well_formed_xml_parts(saved)
        text = words(soffice_convert(saved, "txt:Text (encoded):UTF8", tmp_path).read_text(encoding="utf-8-sig"))
        for t in DOC_TEXTS + ["0,0", "1,1"]:
            assert t in text, f"{fmt}: {t!r} missing in {text[:300]!r}"
        if fmt == "docx":
            d = Document(str(saved))
            assert [p.text for p in d.paragraphs if p.text][:2] == ["Seed document", "Plain, bold and italic."]
            assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "1,1"
            assert d.inline_shapes, "the picture was lost"
        else:
            with zipfile.ZipFile(saved) as z:
                assert z.namelist()[0] == "mimetype" and z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
                assert z.read("mimetype") == b"application/vnd.oasis.opendocument.text"
                assert any(n.startswith("Pictures/") for n in z.namelist()), "the picture was lost"
    # an ODT made by LibreOffice opens too
    odt = soffice_convert(seed(tmp_path, corpus.docx(), "lo.docx"), "odt", tmp_path)
    open_dialog(page, odt)
    expect(page.locator("[data-testid=file-name]")).to_contain_text(odt.name, timeout=30000)
    page.wait_for_timeout(1000)
    saved = save_as(page, "docx", tmp_path)
    text = words(soffice_convert(saved, "txt:Text (encoded):UTF8", tmp_path).read_text(encoding="utf-8-sig"))
    assert all(t in text for t in DOC_TEXTS), text[:300]
    assert page.errors == []


def test_spreadsheet_round_trip(app_page, tmp_path):
    from openpyxl import load_workbook
    page = app_page("spreadsheet", init_script=NO_PICKERS)
    open_dialog(page, seed(tmp_path, corpus.xlsx(), "seed.xlsx"))
    expect(page.locator("[data-testid=file-name]")).to_contain_text("seed.xlsx", timeout=30000)
    page.wait_for_timeout(1000)
    for fmt in ("xlsx", "ods"):
        saved = save_as(page, fmt, tmp_path)
        well_formed_xml_parts(saved)
        import csv as csvlib
        text = soffice_convert(saved, "csv:Text - txt - csv (StarCalc):44,34,76", tmp_path).read_text(encoding="utf-8")
        rows = [[c.replace(",", ".") for c in r] for r in csvlib.reader(io.StringIO(text))]  # (LibreOffice's locale may use a decimal comma)
        assert rows[0][:3] == ["Name", "Value", "Double"], rows[:2]
        assert rows[2][:3] == ["row 2", "3", "6"], rows[:3]  # the formula, recalculated by LibreOffice
        assert rows[0][4] == "22.5", rows[0]  # =SUM(B2:B6)
        assert rows[7][0] == "merged"
        if fmt == "xlsx":
            wb = load_workbook(saved)
            assert wb.sheetnames == ["Data", "Second"]
            ws = wb["Data"]
            assert ws["C2"].value == "=B2*2" and ws["E1"].value == "=SUM(B2:B6)"
            assert "A8:C8" in [str(r) for r in ws.merged_cells.ranges]
            assert ws["A1"].font.bold
    assert page.errors == []


def test_presentation_round_trip(app_page, tmp_path):
    from pptx import Presentation
    page = app_page("presentation", init_script=NO_PICKERS)
    drop_file(page, "seed.pptx", corpus.pptx())
    page.wait_for_function("() => PR.state.name === 'seed.pptx'", timeout=30000)
    saved = download(page, lambda: page.click("[data-testid=save]"), tmp_path)
    well_formed_xml_parts(saved)
    prs = Presentation(str(saved))
    texts = [[sh.text_frame.text for sh in s.shapes if sh.has_text_frame] for s in prs.slides]
    assert len(prs.slides) == 3
    assert "Seed presentation" in texts[0] and "one\ntwo" in texts[1]
    assert prs.slides[1].notes_slide.notes_text_frame.text == "notes"
    assert any(sh.shape_type == 13 for sh in prs.slides[2].shapes), "the picture was lost"
    pdf = soffice_convert(saved, "pdf", tmp_path)
    text = words(run(tool("pdftotext"), pdf, "-").stdout)
    assert "Seed presentation" in text and "Bullets" in text
    assert page.errors == []


def test_web_slides_to_powerpoint(app_page, tmp_path):
    from pptx import Presentation
    page = app_page("reveal", init_script=NO_PICKERS)
    drop_file(page, "seed.html", corpus.REVEAL_HTML.encode())
    page.wait_for_function("() => RV.state.name === 'seed.html'")
    html = download(page, lambda: page.click("[data-testid=save]"), tmp_path)
    r = run(tool("xmllint"), "--html", "--noout", html)
    assert "Opening and ending tag mismatch" not in r.stderr, r.stderr[:500]
    page.click("[data-testid=save-as]")
    page.select_option("[data-testid=save-format]", "pptx")
    saved = download(page, lambda: page.click("[data-testid=save-ok]"), tmp_path)
    well_formed_xml_parts(saved)
    prs = Presentation(str(saved))
    assert "Seed deck" in [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    soffice_convert(saved, "pdf", tmp_path)
    assert page.errors == []


# ---- archives ------------------------------------------------------------------------------------------------------------

def archive_listing(path):
    r = run(tool("7z"), "l", "-slt", "-ba", path)
    assert r.returncode == 0, r.stdout + r.stderr
    return sorted(m for m in re.findall(r"^Path = (.+)$", r.stdout, re.M))


def test_archive_formats(app_page, tmp_path):
    page = app_page("archive", init_script=NO_PICKERS)
    files = {"readme.txt": b"hello\n", "dir/data.bin": bytes(range(256)) * 4, "unicode-äöü.txt": "ü".encode()}
    drop_file(page, "seed.zip", corpus.zip_of(files))
    expect(page.locator("[data-testid=row]").first).to_be_visible()
    for fmt in ("zip", "7z", "tar", "tgz"):
        if fmt == "7z":  # libarchive (WebAssembly, "C" locale) cannot write other names than ASCII: refused with a message
            page.click("[data-testid=save-as]")
            page.select_option("[data-testid=save-as-format]", "7z")
            page.click("[data-testid=save-as-ok]")
            expect(page.locator("[data-testid=toast]").last).to_contain_text("unicode-äöü.txt")
            page.keyboard.press("Escape")
            continue
        saved = save_as(page, fmt, tmp_path)
        assert run(tool("7z"), "t", saved).returncode == 0, f"7-Zip: {saved.name} is broken"
        if fmt == "tgz":  # (7-Zip shows the gzip layer): gzip and Python's tarfile check it
            assert run(tool("gzip"), "-t", saved).returncode == 0
            with tarfile.open(saved) as t:
                assert {m.name: t.extractfile(m).read() for m in t.getmembers() if m.isfile()} == files
            continue
        listing = archive_listing(saved)
        for name in files:
            assert name in listing, f"{fmt}: {name} missing in {listing}"
        out = tmp_path / f"x-{fmt}"
        assert run(tool("7z"), "x", f"-o{out}", saved).returncode == 0
        for name, data in files.items():
            assert (out / name).read_bytes() == data, f"{fmt}: {name}"
        if fmt == "zip":
            assert run(tool("unzip"), "-tq", saved).returncode == 0
    assert page.errors == []


# ---- diagrams, images ----------------------------------------------------------------------------------------------------

def test_diagram_formats(app_page, tmp_path):
    from lxml import etree
    from PIL import Image
    page = app_page("diagram", init_script=NO_PICKERS)
    drop_file(page, "seed.drawio", corpus.DRAWIO.encode())
    page.wait_for_function("() => document.title.startsWith('seed.drawio')")
    drawio = download(page, lambda: page.click("[data-testid=save]"), tmp_path)
    root = etree.parse(str(drawio)).getroot()
    assert root.tag == "mxfile"
    import html as htmllib
    text = lambda c: htmllib.unescape(re.sub(r"<[^>]+>", "", c.get("value") or "")) if "html=1" in (c.get("style") or "") else c.get("value")
    cells = [c for c in root.iter("mxCell") if c.get("vertex") or c.get("edge")]
    assert sorted(text(c) for c in cells if text(c)) == ["End", "Start & more", "goes", "in group"]  # (labels are plain text, documented)
    group = next(c for c in cells if c.get("style") == "group")
    assert next(c for c in cells if text(c) == "in group").get("parent") == group.get("id")  # still in the group
    edge = next(c for c in cells if c.get("edge"))
    assert [(p.get("x"), p.get("y")) for p in edge.iter("mxPoint")] == [("200", "20")]  # the routing point
    svg = download(page, lambda: page.click("[data-testid=export-svg]"), tmp_path)
    assert run(tool("xmllint"), "--noout", svg).returncode == 0
    png = download(page, lambda: page.click("[data-testid=export-png]"), tmp_path)
    img = Image.open(png)
    img.verify()
    assert "mxfile" in Image.open(png).info  # editable PNG, as draw.io writes it
    assert page.errors == []


@pytest.mark.parametrize("fmt", ["png", "jpeg", "webp", "bmp", "ora", "psd"])
def test_image_formats(app_page, tmp_path, fmt):
    from PIL import Image
    page = app_page("imageeditor", init_script=NO_PICKERS)
    drop_file(page, "seed.png", corpus.image("PNG", (40, 30)), "image/png")
    page.wait_for_function("() => document.title.startsWith('seed.png')")
    page.click("[data-testid=menu-file] summary")
    page.click("[data-testid=cmd-saveAs]")
    page.select_option("[data-testid=dlg-format]", fmt)
    saved = download(page, lambda: page.click("[data-testid=dlg-ok]"), tmp_path)
    if fmt == "ora":
        with zipfile.ZipFile(saved) as z:
            assert z.namelist()[0] == "mimetype" and z.read("mimetype") == b"image/openraster"
            assert z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
            stack = z.read("stack.xml")
            assert b"<image" in stack
            img = Image.open(io.BytesIO(z.read("mergedimage.png")))
    elif fmt == "psd":
        from psd_tools import PSDImage
        img = PSDImage.open(saved).composite()
    else:
        img = Image.open(saved)
        img.load()
        r = run(tool("ffprobe"), "-v", "error", "-show_entries", "stream=width,height", "-of", "csv=p=0", saved)
        assert r.stdout.strip() == "40,30", r.stdout
    assert img.size == (40, 30)
    px = img.convert("RGB").getpixel((20, 2))
    assert abs(px[0] - 30) < 12 and abs(px[2] - 200) < 12, px  # the colour of the seed (JPEG/WebP: lossy)
    assert page.errors == []


# ---- data ----------------------------------------------------------------------------------------------------------------

def test_calendar_round_trip(app_page, tmp_path):
    import icalendar
    page = app_page("calendar", init_script=NO_PICKERS)
    drop_file(page, "seed.ics", corpus.ICS.encode(), "text/calendar")
    expect(page.locator("[data-testid=calendar]")).not_to_have_count(0)
    page.wait_for_timeout(500)
    saved = download(page, lambda: page.locator("[data-testid=cal-save]").last.click(), tmp_path)
    cal = icalendar.Calendar.from_ical(saved.read_bytes())
    events = {str(e["UID"]): e for e in cal.walk("VEVENT")}
    weekly = events["seed-1@example.org"]
    assert str(weekly["SUMMARY"]) == "Weekly meeting"
    assert str(weekly["DESCRIPTION"]) == "Line one\nLine two, with comma"
    assert weekly["RRULE"]["FREQ"] == ["WEEKLY"] and weekly["RRULE"]["COUNT"] == [5]
    assert weekly["DTSTART"].dt.tzinfo is not None and weekly["DTSTART"].dt.hour == 10
    assert weekly.get("EXDATE"), "the exception date was lost"
    assert list(weekly.walk("VALARM")), "the reminder was lost"
    assert str(events["seed-2@example.org"]["SUMMARY"]) == "All day"
    assert [str(t["SUMMARY"]) for t in cal.walk("VTODO")] == ["A task"]
    assert page.errors == []


def test_keepass_round_trip(app_page, tmp_path):
    from pykeepass import PyKeePass
    page = app_page("keepass", init_script=NO_PICKERS)
    drop_file(page, "seed.kdbx", corpus.kdbx(tmp_path))
    page.fill("[data-testid=password]", "pw")
    page.keyboard.press("Enter")
    expect(page.locator("[data-testid=group]").first).to_be_visible(timeout=30000)
    saved = download(page, lambda: page.click("[data-testid=save]"), tmp_path)
    kp = PyKeePass(str(saved), password="pw")
    entry = kp.find_entries(title="Example", first=True)
    assert entry and entry.username == "ann" and entry.password == "secret" and entry.url == "https://example.org"
    assert entry.notes == "line\nline" and entry.get_custom_property("extra") == "value"
    assert entry.group.name == "Web"
    assert kp.find_entries(title="Root entry", first=True)
    assert page.errors == []


def test_database_round_trip(app_page, tmp_path):
    page = app_page("dbtool", init_script=NO_PICKERS)
    drop_file(page, "seed.sqlite", corpus.sqlite_db(tmp_path))
    expect(page.locator("[data-testid=table]").first).to_be_visible(timeout=30000)
    saved = download(page, lambda: page.click("[data-testid=save-db]"), tmp_path)
    con = sqlite3.connect(saved)
    assert con.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    assert con.execute("SELECT name, born, hex(photo) FROM people ORDER BY id").fetchall() == [("Ann", "1990-01-01", "89504E47"), ("Bob", None, "")]  # (hex(NULL) is "")
    names = {r[0] for r in con.execute("SELECT name FROM sqlite_master")}
    assert {"people", "orders", "people_name", "adults", "t"} <= names
    con.close()
    assert page.errors == []


def test_notebook_round_trip(app_page, tmp_path):
    import nbformat
    page = app_page("notebook", init_script=NO_PICKERS)
    drop_file(page, "seed.ipynb", json.dumps(corpus.NOTEBOOK).encode())
    expect(page.locator("[data-testid=output]").first).to_be_visible(timeout=30000)
    saved = download(page, lambda: page.keyboard.press("Control+s"), tmp_path)
    nb = nbformat.read(str(saved), as_version=4)
    nbformat.validate(nb)
    assert [c.cell_type for c in nb.cells] == ["markdown", "code", "raw"]
    assert nb.cells[0].source == "# Seed\nSome *Markdown* and $x^2$."
    assert nb.cells[1].metadata["slideshow"]["slide_type"] == "slide"
    outputs = nb.cells[1].outputs
    assert outputs[0].data["text/plain"] == "2" and "image/png" in outputs[2].data
    assert page.errors == []


def test_pdf_round_trip(app_page, tmp_path):
    from pypdf import PdfReader
    page = app_page("pdfviewer", init_script=NO_PICKERS)
    drop_file(page, "seed.pdf", corpus.minimal_pdf())
    page.wait_for_selector(".page[data-page-number='1'] .textLayer span", timeout=30000)
    saved = download(page, lambda: page.click("[data-testid=save]"), tmp_path)
    info = run(tool("pdfinfo"), saved)
    assert info.returncode == 0 and re.search(r"Pages:\s+1", info.stdout), info.stdout + info.stderr
    assert "Seed PDF" in run(tool("pdftotext"), saved, "-").stdout
    reader = PdfReader(str(saved))
    assert reader.outline and reader.outline[0].title == "Start"
    assert page.errors == []


def test_audio_round_trip(app_page, tmp_path):
    page = app_page("audioeditor", init_script=NO_PICKERS)
    src = seed(tmp_path, corpus.wav(seconds=1.0, rate=44100), "seed.wav")
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(src))
    expect(page.locator("[data-testid=track]")).to_have_count(1)
    page.click("[data-testid=export]")
    page.select_option("[data-testid=export-format]", "wav16")
    saved = download(page, lambda: page.click("[data-testid=export-ok]"), tmp_path)
    probe = run(tool("ffprobe"), "-v", "error", "-show_entries", "stream=codec_name,sample_rate,channels,duration", "-of", "json", saved)
    stream = json.loads(probe.stdout)["streams"][0]
    assert stream["codec_name"].startswith("pcm_") and int(stream["sample_rate"]) == 44100
    assert abs(float(stream["duration"]) - 1.0) < 0.01
    with wave.open(str(saved)) as w, wave.open(str(src)) as orig:
        assert w.getnframes() == orig.getnframes()
    assert page.errors == []
