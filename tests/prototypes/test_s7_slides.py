"""S-7: presentation editor (shared/js/pptx-import.js, slide-render.js, pptx-export.js) and Univer Slides."""
import base64
import io
import re

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.util import Emu

from helpers import office_samples
from helpers.paths import dist_url


@pytest.fixture(scope="session")
def sample_pptx(tmp_path_factory):
    path = tmp_path_factory.mktemp("pptx") / "sample.pptx"
    office_samples.make_pptx(path)
    return path


def open_editor(page, sample=None):
    page.goto(dist_url("s7-slides"))
    page.wait_for_selector("body[data-ready]", timeout=30000)
    if sample:
        page.evaluate("([b, n]) => S7.openBase64(b, n)", [base64.b64encode(sample.read_bytes()).decode(), sample.name])


def deck(page):
    return page.evaluate("S7.state.deck")


def texts(element):
    return ["".join(r["text"] for r in p["runs"]) for p in element["text"]["paragraphs"]]


# ---- import ----------------------------------------------------------------------------------------

def test_import(page, sample_pptx):
    open_editor(page, sample_pptx)
    d = deck(page)
    assert (d["width"], d["height"], len(d["slides"])) == (960, 720, 3)
    # slide 1: placeholders without own position inherit position, size, alignment and font size
    title, subtitle = d["slides"][0]["elements"]
    assert texts(title) == ["Quartalsbericht 2026"] and texts(subtitle) == ["Vertraulich – äöü"]
    assert [round(title[k]) for k in ("x", "y", "w", "h")] == [72, 224, 816, 154]
    assert title["text"]["paragraphs"][0]["align"] == "ctr" and title["text"]["paragraphs"][0]["runs"][0]["size"] == 44
    assert subtitle["text"]["paragraphs"][0]["bullet"] is None  # the layout switches the master's bullets off
    # slide 2: bullets and sizes per level from the master's body style
    body = d["slides"][1]["elements"][1]
    paragraphs = body["text"]["paragraphs"]
    assert texts(body) == ["Umsatz", "Inland", "Export", "Kosten"]
    assert [p["level"] for p in paragraphs] == [0, 1, 1, 0]
    assert [p["bullet"]["char"] for p in paragraphs] == ["•", "–", "–", "•"]
    assert [p["runs"][0]["size"] for p in paragraphs] == [32, 28, 28, 32]
    # slide 3: background, run formatting, shapes, image, table
    s3 = d["slides"][2]
    assert s3["background"] == "#FFF2CC"
    box, rect, oval, image, table = s3["elements"]
    runs = {r["text"]: r for r in box["text"]["paragraphs"][0]["runs"]}
    assert runs["fett"]["bold"] and runs["kursiv rot"]["italic"]
    assert (runs["kursiv rot"]["color"], runs["kursiv rot"]["size"]) == ("#FF0000", 28)
    assert (rect["type"], rect["geometry"], rect["fill"]) == ("shape", "rect", "#4472C4")
    assert rect["text"]["paragraphs"][0]["runs"][0]["color"] == "#FFFFFF"  # from the shape style (fontRef lt1)
    assert (oval["geometry"], oval["fill"], oval["line"]["color"], oval["line"]["width"]) == ("ellipse", None, "#008000", 4)
    assert image["type"] == "image" and image["image"]["src"].startswith("data:image/png;base64,")
    assert [[texts(c)[0] for c in r["cells"]] for r in table["table"]["rows"]] == [["Name", "Wert"], ["Äpfel", "3"]]
    assert page.errors == [] and page.requests == []


# ---- export ----------------------------------------------------------------------------------------

def test_round_trip(page, sample_pptx, tmp_path):
    open_editor(page, sample_pptx)
    out = base64.b64decode(page.evaluate("S7.exportBase64()"))
    prs = Presentation(io.BytesIO(out))
    assert (prs.slide_width, prs.slide_height, len(prs.slides)) == (9144000, 6858000, 3)
    s1, s2, s3 = prs.slides
    title = s1.shapes[0]
    assert title.text_frame.text == "Quartalsbericht 2026" and title.text_frame.paragraphs[0].runs[0].font.size.pt == 44
    assert round(Emu(title.left).inches, 2) == 0.75 and round(Emu(title.top).inches, 2) == 2.33
    body = s2.shapes[1].text_frame
    assert [p.text for p in body.paragraphs] == ["Umsatz", "Inland", "Export", "Kosten"]
    assert [p.level for p in body.paragraphs] == [0, 1, 1, 0]
    shapes = list(s3.shapes)
    runs = {r.text: r for r in shapes[0].text_frame.paragraphs[0].runs}
    assert runs["fett"].font.bold and runs["kursiv rot"].font.italic
    assert (str(runs["kursiv rot"].font.color.rgb), runs["kursiv rot"].font.size.pt) == ("FF0000", 28)
    assert shapes[1].auto_shape_type == MSO_AUTO_SHAPE_TYPE.RECTANGLE and str(shapes[1].fill.fore_color.rgb) == "4472C4"
    assert shapes[2].auto_shape_type == MSO_AUTO_SHAPE_TYPE.OVAL and str(shapes[2].line.color.rgb) == "008000"
    assert shapes[3].shape_type == MSO_SHAPE_TYPE.PICTURE
    assert [[c.text for c in r.cells] for r in shapes[4].table.rows] == [["Name", "Wert"], ["Äpfel", "3"]]
    if office_samples.SOFFICE:  # LibreOffice must accept the file
        path = tmp_path / "export.pptx"
        path.write_bytes(out)
        pdf = office_samples.convert(path, "pdf", tmp_path / "pdf")
        assert len(re.findall(rb"/Type\s*/Page(?!s)", pdf.read_bytes())) == 3
    assert page.errors == []


# ---- editor ----------------------------------------------------------------------------------------

def stage_scale(page):
    return page.evaluate("document.getElementById('stage').getBoundingClientRect().width / S7.state.deck.width")


def test_move_and_resize(page, sample_pptx):
    open_editor(page, sample_pptx)
    page.click("#thumbs .thumb:nth-child(3)")
    rect = page.locator('#stage .el[data-id="el7"]')  # the blue rectangle
    before = page.evaluate("S7.state.deck.slides[2].elements.find(e => e.id === 'el7')")
    box = rect.bounding_box()
    page.mouse.move(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
    page.mouse.down()
    page.mouse.move(box["x"] + box["width"] / 2 + 100, box["y"] + box["height"] / 2 + 50, steps=5)
    page.mouse.up()
    after = page.evaluate("S7.state.deck.slides[2].elements.find(e => e.id === 'el7')")
    s = stage_scale(page)
    assert after["x"] - before["x"] == pytest.approx(100 / s, abs=2)
    assert after["y"] - before["y"] == pytest.approx(50 / s, abs=2)
    # resize with the bottom right handle
    handle = page.locator(".moveable-control.moveable-se")
    hb = handle.bounding_box()
    page.mouse.move(hb["x"] + hb["width"] / 2, hb["y"] + hb["height"] / 2)
    page.mouse.down()
    page.mouse.move(hb["x"] + 60, hb["y"] + 30, steps=5)
    page.mouse.up()
    resized = page.evaluate("S7.state.deck.slides[2].elements.find(e => e.id === 'el7')")
    assert resized["w"] > after["w"] + 40 / s and resized["h"] > after["h"] + 15 / s
    assert page.errors == []


def test_edit_text_and_format(page):
    open_editor(page)
    title = page.locator('#stage .el[data-id="title"]')
    title.dblclick()
    page.keyboard.press("Control+a")
    page.keyboard.type("Neuer Titel mit Ümlaut")
    page.click("#stage-wrap", position={"x": 5, "y": 5})  # blur: the text is read back into the model
    el = page.evaluate("S7.state.deck.slides[0].elements[0]")
    assert texts(el) == ["Neuer Titel mit Ümlaut"]
    assert el["text"]["paragraphs"][0]["runs"][0]["size"] == 40
    page.click('#stage .el[data-id="title"]')
    page.click("#italic")
    assert all(r["italic"] for r in page.evaluate("S7.state.deck.slides[0].elements[0]")["text"]["paragraphs"][0]["runs"])
    assert page.errors == []


def test_add_and_delete(page):
    open_editor(page)
    page.click("#add-slide")
    page.click("#add-rect")
    page.click("#add-ellipse")
    d = deck(page)
    assert len(d["slides"]) == 2 and [e["geometry"] for e in d["slides"][1]["elements"]] == ["rect", "ellipse"]
    page.keyboard.press("Delete")  # the ellipse is selected after adding it
    assert [e["geometry"] for e in deck(page)["slides"][1]["elements"]] == ["rect"]
    assert page.locator("#thumbs .thumb").count() == 2


def test_present(page, sample_pptx):
    open_editor(page, sample_pptx)
    page.click("#thumbs .thumb:nth-child(1)")
    page.click("#present")
    presenter = page.locator("#presenter")
    assert presenter.is_visible() and presenter.get_attribute("data-index") == "0"
    page.keyboard.press("ArrowRight")
    page.keyboard.press("ArrowRight")
    assert presenter.get_attribute("data-index") == "2"
    page.keyboard.press("ArrowLeft")
    assert presenter.get_attribute("data-index") == "1"
    page.keyboard.press("Escape")
    assert presenter.is_hidden()


def test_ui_open_and_save(page, sample_pptx, tmp_path):
    open_editor(page)
    page.set_input_files("#open", str(sample_pptx))
    page.wait_for_selector("#status:has-text('3 slides')")
    assert page.locator("#thumbs .thumb").count() == 3
    with page.expect_download() as download:
        page.click("#save")
    assert download.value.suggested_filename == "sample.pptx"
    path = tmp_path / "saved.pptx"
    download.value.save_as(path)
    assert len(Presentation(path).slides) == 3


# ---- Univer Slides (evaluated, not chosen) ---------------------------------------------------------

def test_univer_slides_loads(page):
    page.goto(dist_url("s7-univer-slides"))
    page.wait_for_selector("body[data-ready]", timeout=30000)
    page.wait_for_timeout(2000)
    assert page.locator("canvas").count() >= 1
    # Univer tries to load hyphenation patterns for the detected language (not vendored); Firefox logs
    # that error as an object ("JSHandle@object")
    assert [e for e in page.errors if "@univerjs/engine-render/lib/es/" not in e and e != "JSHandle@object"] == []
