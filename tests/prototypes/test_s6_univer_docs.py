"""S-6: DOCX import and export for Univer Docs (shared/js/univer-docx.js; since S-8 with the direct DOCX reader)."""
import base64
import io

import docx
import pytest

from helpers import office_samples
from helpers.paths import dist_url


@pytest.fixture(scope="session")
def sample_docx(tmp_path_factory):
    return office_samples.make_docx(tmp_path_factory.mktemp("docx") / "sample.docx")


def open_page(page):
    page.goto(dist_url("s6-univer-docs"))
    page.wait_for_selector("body[data-ready]", timeout=60000)


def import_docx(page, path):
    result = page.evaluate("b => S6D.importBase64(b)", base64.b64encode(path.read_bytes()).decode())
    page.wait_for_timeout(1000)
    return result


def snapshot(page):
    return page.evaluate("S6D.univerAPI.getActiveDocument().save()")


def export(page):
    return base64.b64decode(page.evaluate("S6D.exportBase64()"))


def test_import(page, sample_docx):
    open_page(page)
    import_docx(page, sample_docx)
    snap = snapshot(page)
    body = snap["body"]
    text = body["dataStream"]
    for expected in ("Projektbericht", "Normal text with bold, italic, underlined and red.", "Erster Punkt", "Schritt zwei",
                     "Äpfel", "Zentriert äöü ß 🍒", "Second page"):
        assert expected in text
    runs = {text[r["st"]:r["ed"]]: r["ts"] for r in body["textRuns"]}
    assert runs["bold"]["bl"] == 1 and runs["italic"]["it"] == 1 and runs["underlined"]["ul"] == {"s": 1}
    assert "bl" not in runs["italic"] and "it" not in runs["bold"]
    styles = [p.get("paragraphStyle", {}).get("namedStyleType") for p in body["paragraphs"]]
    assert styles[0] == 4 and 5 in styles  # heading 1 and 2
    lists = [p["bullet"]["listType"] for p in body["paragraphs"] if p.get("bullet")]
    assert lists == ["BULLET_LIST"] * 3 + ["ORDER_LIST"] * 2
    assert len(body["tables"]) == 1 and len(snap["tableSource"]) == 1
    assert snap["documentStyle"]["documentFlavor"] == 1 and snap["documentStyle"]["autoHyphenation"] == 0
    assert page.errors == [] and page.requests == []


def test_round_trip(page, sample_docx, tmp_path):
    open_page(page)
    import_docx(page, sample_docx)
    out = export(page)
    d = docx.Document(io.BytesIO(out))
    paragraphs = [(p.style.name, p.text) for p in d.paragraphs]
    assert paragraphs[0] == ("Heading 1", "Projektbericht")
    assert ("Heading 2", "Liste") in paragraphs
    assert [t for s, t in paragraphs if s == "List Paragraph"] == ["Erster Punkt", "Zweiter Punkt", "Dritter Punkt", "Schritt eins", "Schritt zwei"]
    runs = {r.text: (r.bold, r.italic, r.underline) for p in d.paragraphs for r in p.runs}
    assert runs["bold"] == (True, None, None) and runs["italic"] == (None, True, None) and runs["underlined"] == (None, None, True)
    numbered = [p for p in d.paragraphs if p.text.startswith("Schritt")]
    assert all(p._p.pPr.numPr is not None for p in numbered)
    assert [[c.text for c in row.cells] for row in d.tables[0].rows] == [["Name", "Menge", "Preis"], ["Äpfel", "3", "1,50"]]
    if office_samples.SOFFICE:  # LibreOffice must accept the file
        path = tmp_path / "export.docx"
        path.write_bytes(out)
        assert office_samples.convert(path, "odt", tmp_path / "odt").stat().st_size > 0
    assert page.errors == []


def test_typing_and_export(page):
    open_page(page)
    page.click("canvas", position={"x": 400, "y": 120})
    page.keyboard.type("Typed in the browser")
    page.wait_for_timeout(500)
    assert "Typed in the browser" in snapshot(page)["body"]["dataStream"]
    assert any(p.text == "Typed in the browser" for p in docx.Document(io.BytesIO(export(page))).paragraphs)


def test_ui_open_and_save(page, sample_docx, tmp_path):
    open_page(page)
    page.set_input_files("#open", str(sample_docx))
    page.wait_for_selector("#status:has-text('sample.docx')", timeout=30000)
    with page.expect_download() as download:
        page.click("#export-docx")
    assert download.value.suggested_filename == "sample.docx"
    path = tmp_path / "saved.docx"
    download.value.save_as(path)
    assert docx.Document(path).paragraphs[0].text == "Projektbericht"
