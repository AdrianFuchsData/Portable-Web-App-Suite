"""S-8: direct DOCX reader (shared/js/docx-import.js) and the extended DOCX export (shared/js/univer-docx.js)."""
import base64
import io
import re
import zipfile

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Emu

from helpers import office_samples
from helpers.paths import dist_url


@pytest.fixture(scope="session")
def rich_docx(tmp_path_factory):
    return office_samples.make_docx_rich(tmp_path_factory.mktemp("rich") / "rich.docx")


def open_page(page):
    page.goto(dist_url("s6-univer-docs"))
    page.wait_for_selector("body[data-ready]", timeout=60000)


def import_docx(page, path, via_html=False):
    result = page.evaluate("([b, h]) => S6D.importBase64(b, h)", [base64.b64encode(path.read_bytes()).decode(), via_html])
    page.wait_for_timeout(1000)
    return result


def snapshot(page):
    return page.evaluate("S6D.univerAPI.getActiveDocument().save()")


def runs_by_text(body):
    s = body["dataStream"]
    return {s[r["st"]:r["ed"]]: r["ts"] for r in body["textRuns"]}


def paragraph_of(body, text):
    s = body["dataStream"]
    start = s.index(text)
    return next(p for p in body["paragraphs"] if p["startIndex"] >= start)


def test_direct_import(page, rich_docx):
    open_page(page)
    assert import_docx(page, rich_docx)["messages"] == []
    snap = snapshot(page)
    body = snap["body"]
    runs = runs_by_text(body)
    # run formatting incl. style inheritance (theme fonts, heading style with own color and size)
    assert runs["Einleitung"]["cl"]["rgb"] == "#1F3864" and runs["Einleitung"]["fs"] == 20 and runs["Einleitung"]["bl"] == 1
    assert (runs["Rot 14pt "]["cl"]["rgb"], runs["Rot 14pt "]["fs"]) == ("#FF0000", 14)
    assert runs["Arial "]["ff"] == "Arial" and runs["Rot 14pt "]["ff"] == "Cambria"  # theme minor font
    assert runs["markiert"]["bg"]["rgb"] == "#FFFF00"
    assert runs["2"]["va"] == 3 and runs[" durchgestrichen"]["st"]["s"] == 1
    # paragraph formatting
    assert paragraph_of(body, "Dokumenttitel")["paragraphStyle"]["namedStyleType"] == 2
    assert paragraph_of(body, "Einleitung")["paragraphStyle"]["namedStyleType"] == 4
    assert paragraph_of(body, "Listen")["paragraphStyle"]["namedStyleType"] == 5
    centered = paragraph_of(body, "Zentriert")["paragraphStyle"]
    assert (centered["horizontalAlign"], centered["spaceAbove"]["v"], centered["spaceBelow"]["v"]) == (2, 12, 18)
    assert paragraph_of(body, "Rechtsbündig")["paragraphStyle"]["horizontalAlign"] == 3
    justified = paragraph_of(body, "Blocksatz")["paragraphStyle"]
    assert (justified["horizontalAlign"], justified["lineSpacing"]) == (4, 1.5)
    assert justified["indentStart"]["v"] == pytest.approx(28.35, abs=0.1)  # 1 cm
    lists = [(p["bullet"]["listType"]) for p in body["paragraphs"] if p.get("bullet")]
    assert lists == ["BULLET_LIST"] * 3 + ["ORDER_LIST"] * 2
    # table: shading and horizontal merge
    table = list(snap["tableSource"].values())[0]
    assert [c.get("backgroundColor", {}).get("rgb") for c in table["tableRows"][0]["tableCells"]] == ["#D9E2F3"] * 3
    assert table["tableRows"][2]["tableCells"][0]["columnSpan"] == 2
    # image, hyperlink, page break
    assert len(body["customBlocks"]) == 1
    drawing = snap["drawings"][body["customBlocks"][0]["blockId"]]
    assert drawing["source"].startswith("data:image/png;base64,")
    assert drawing["docTransform"]["size"]["width"] == pytest.approx(6 / 2.54 * 96, abs=1)
    link = body["customRanges"][0]
    assert link["properties"]["url"] == "https://example.org/info"
    assert body["dataStream"][link["startIndex"]:link["endIndex"] + 1] == "example.org"
    assert "\f" in body["dataStream"]
    # page setup, header and footer
    ds = snap["documentStyle"]
    assert ds["pageSize"]["width"] == pytest.approx(21 / 2.54 * 96, abs=0.5) and ds["marginLeft"] == pytest.approx(2 / 2.54 * 96, abs=0.5)
    assert snap["headers"][ds["defaultHeaderId"]]["body"]["dataStream"].startswith("Kopfzeile Projekt X")
    assert snap["footers"][ds["defaultFooterId"]]["body"]["dataStream"].startswith("Fußzeile – vertraulich")
    page.wait_for_selector("text=/(Seite|Page) 1 (von|of) 2/", timeout=20000)  # two pages (counted after layout)
    assert page.errors == [] and page.requests == []


def test_mammoth_route_loses_formatting(page, rich_docx):
    """For comparison: the HTML route (S-6) drops colors, sizes, alignment, headers and footers."""
    open_page(page)
    import_docx(page, rich_docx, via_html=True)
    snap = snapshot(page)
    runs = runs_by_text(snap["body"])
    assert "cl" not in runs.get("Rot 14pt ", {}) and "fs" not in runs.get("Rot 14pt ", {})
    assert "horizontalAlign" not in paragraph_of(snap["body"], "Zentriert").get("paragraphStyle", {})
    assert not snap["documentStyle"].get("defaultHeaderId")


def test_round_trip(page, rich_docx, tmp_path):
    open_page(page)
    import_docx(page, rich_docx)
    out = base64.b64decode(page.evaluate("S6D.exportBase64()"))
    d = docx.Document(io.BytesIO(out))
    section = d.sections[0]
    assert (round(Emu(section.page_width).cm, 1), round(Emu(section.page_height).cm, 1)) == (21.0, 29.7)
    assert round(Emu(section.left_margin).cm, 1) == 2.0
    assert section.header.paragraphs[0].text == "Kopfzeile Projekt X"
    assert section.footer.paragraphs[0].text == "Fußzeile – vertraulich"
    paragraphs = {p.text: p for p in d.paragraphs if p.text}
    assert paragraphs["Dokumenttitel"].style.name == "Title"
    assert paragraphs["Einleitung"].style.name == "Heading 1" and paragraphs["Listen"].style.name == "Heading 2"
    heading_run = paragraphs["Einleitung"].runs[0]
    assert (str(heading_run.font.color.rgb), heading_run.font.size.pt) == ("1F3864", 20)
    runs = {r.text: r for r in paragraphs["Rot 14pt Arial markiert E=mc2 durchgestrichen"].runs}
    assert (str(runs["Rot 14pt "].font.color.rgb), runs["Rot 14pt "].font.size.pt, runs["Arial "].font.name) == ("FF0000", 14, "Arial")
    assert runs["markiert"].font.highlight_color == WD_COLOR_INDEX.YELLOW
    assert runs["2"].font.superscript and runs[" durchgestrichen"].font.strike
    centered = paragraphs["Zentriert mit Abstand"]
    assert centered.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert (centered.paragraph_format.space_before.pt, centered.paragraph_format.space_after.pt) == (12, 18)
    assert paragraphs["Rechtsbündig"].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    justified = next(p for t, p in paragraphs.items() if t.startswith("Blocksatz"))
    assert justified.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY and justified.paragraph_format.line_spacing == 1.5
    assert round(Emu(justified.paragraph_format.left_indent).cm, 1) == 1.0
    assert all(paragraphs[t]._p.pPr.numPr is not None for t in ("Punkt eins", "Schritt zwei"))
    rows = [[c.text for c in r.cells] for r in d.tables[0].rows]
    assert rows == [["Name", "Menge", "Preis"], ["Äpfel", "3", "1,50"], ["Summe", "Summe", "4,50"]]  # merged cell
    xml = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml").decode()
    assert xml.count('w:fill="D9E2F3"') == 3 and 'w:type="page"' in xml
    assert len(d.inline_shapes) == 1 and round(Emu(d.inline_shapes[0].width).cm, 1) == 6.0
    assert [r.target_ref for r in d.part.rels.values() if "hyperlink" in r.reltype] == ["https://example.org/info"]
    if office_samples.SOFFICE:  # LibreOffice must accept the file; the page break gives two pages
        path = tmp_path / "export.docx"
        path.write_bytes(out)
        pdf = office_samples.convert(path, "pdf", tmp_path / "pdf").read_bytes()
        assert len(re.findall(rb"/Type\s*/Page(?!s)", pdf)) == 2
    assert page.errors == []


@pytest.mark.skipif(not office_samples.SOFFICE, reason="LibreOffice not available")
def test_import_libreoffice_docx(page, rich_docx, tmp_path):
    """A DOCX written by LibreOffice (different XML conventions than Word/python-docx)."""
    odt = office_samples.convert(rich_docx, "odt", tmp_path / "odt")
    lo_docx = office_samples.convert(odt, "docx:MS Word 2007 XML", tmp_path / "docx")
    open_page(page)
    import_docx(page, lo_docx)
    snap = snapshot(page)
    body = snap["body"]
    for text in ("Dokumenttitel", "Einleitung", "Punkt eins", "Äpfel", "example.org", "Zweite Seite"):
        assert text in body["dataStream"]
    assert runs_by_text(body)["Rot 14pt "]["cl"]["rgb"] == "#FF0000"
    assert len(body["customBlocks"]) == 1 and len(snap["tableSource"]) == 1
    assert snap["headers"][snap["documentStyle"]["defaultHeaderId"]]["body"]["dataStream"].startswith("Kopfzeile")
    assert page.errors == []
