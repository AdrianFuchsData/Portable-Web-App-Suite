"""S-6: XLSX/ODS/CSV import and export for Univer Sheets (shared/js/univer-xlsx.js, ExcelJS, SheetJS)."""
import base64
import datetime
import io

import openpyxl
import pytest

from helpers import office_samples
from helpers.paths import dist_url

needs_soffice = pytest.mark.skipif(not office_samples.SOFFICE, reason="LibreOffice not available")


@pytest.fixture(scope="session")
def samples(tmp_path_factory):
    d = tmp_path_factory.mktemp("office")
    xlsx = office_samples.make_xlsx(d / "sample.xlsx")
    result = {"xlsx": xlsx}
    if office_samples.SOFFICE:
        result["ods"] = office_samples.convert(xlsx, "ods", d / "ods")
        result["lo_xlsx"] = office_samples.convert(xlsx, "xlsx:Calc MS Excel 2007 XML", d / "lo")
    return result


def open_page(page):
    page.goto(dist_url("s6-office-io"))
    page.wait_for_selector("body[data-ready]", timeout=60000)


def import_file(page, path):
    page.evaluate("([b, n]) => S6.importBase64(b, n)", [base64.b64encode(path.read_bytes()).decode(), path.name])
    page.wait_for_timeout(1500)  # formulas are calculated asynchronously


def value(page, a1, sheet=None):
    return page.evaluate("""([a1, sheet]) => {
      const wb = S6.univerAPI.getActiveWorkbook();
      const ws = sheet ? wb.getSheetByName(sheet) : wb.getActiveSheet();
      return ws.getRange(a1).getValue();
    }""", [a1, sheet])


def export(page, fmt):
    return base64.b64decode(page.evaluate("f => S6.exportBase64(f)", fmt))


def snapshot(page):
    return page.evaluate("S6.univerAPI.getActiveWorkbook().save()")


# ---- import ----------------------------------------------------------------------------------------

def test_import_xlsx(page, samples):
    open_page(page)
    import_file(page, samples["xlsx"])
    snap = snapshot(page)
    assert [snap["sheets"][i]["name"] for i in snap["sheetOrder"]] == ["Data", "Second"]
    data = snap["sheets"][snap["sheetOrder"][0]]
    cells = data["cellData"]
    assert cells["1"]["0"]["v"] == "Äpfel" and cells["3"]["0"]["v"] == "Kirschen 🍒"
    assert cells["1"]["3"]["f"] == "=B2*C2" and cells["2"]["3"]["si"] == cells["1"]["3"]["si"]  # shared formula
    assert cells["1"]["4"]["v"] == 46037  # 2026-01-15 as serial number
    # Univer calculates all formulas (the openpyxl file has no cached results)
    assert [value(page, a) for a in ("D2", "D3", "D4", "D5")] == [4.5, 9, 50, 63.5]
    assert value(page, "A1", "Second") == 127  # cross-sheet formula
    styles = snap["styles"]
    a1 = styles[cells["0"]["0"]["s"]]
    assert a1["bl"] == 1 and a1["cl"]["rgb"] == "#0000FF" and a1["bg"]["rgb"] == "#FFFF00"
    assert a1["bd"]["b"]["s"] == 1 and a1["ht"] == 2 and a1["vt"] == 1
    d5 = styles[cells["4"]["3"]["s"]]
    assert d5["bd"]["t"]["s"] == 7 and d5["ul"]["s"] == 1 and d5["st"]["s"] == 1
    assert styles[cells["1"]["2"]["s"]]["n"]["pattern"] == "#,##0.00"
    assert styles[cells["6"]["0"]["s"]] == {"ff": "Arial", "fs": 16, "ht": 2, "tb": 3}
    assert data["mergeData"] == [{"startRow": 6, "startColumn": 0, "endRow": 6, "endColumn": 3}]
    assert data["columnData"]["0"]["w"] == 145 and data["rowData"]["0"]["h"] == 40
    assert data["freeze"]["ySplit"] == 1
    assert snap["sheets"][snap["sheetOrder"][1]]["tabColor"] == "#00B050"
    assert page.errors == [] and page.requests == []


@needs_soffice
def test_import_libreoffice_xlsx(page, samples):
    """An XLSX written by LibreOffice (theme part, its own style structure, cached results)."""
    open_page(page)
    import_file(page, samples["lo_xlsx"])
    assert value(page, "A2") == "Äpfel"
    assert [value(page, a) for a in ("D2", "D5")] == [4.5, 63.5]
    snap = snapshot(page)
    cells = snap["sheets"][snap["sheetOrder"][0]]["cellData"]
    assert snap["styles"][cells["0"]["0"]["s"]]["bl"] == 1
    assert page.errors == []


@needs_soffice
def test_import_ods(page, samples):
    """ODS goes through SheetJS: values, formulas and merged cells survive, styles do not (SheetJS CE)."""
    open_page(page)
    import_file(page, samples["ods"])
    assert [value(page, a) for a in ("A2", "A4")] == ["Äpfel", "Kirschen 🍒"]
    assert [value(page, a) for a in ("D2", "D3", "D5")] == [4.5, 9, 63.5]
    snap = snapshot(page)
    data = snap["sheets"][snap["sheetOrder"][0]]
    assert data["cellData"]["4"]["3"]["f"] == "=SUM(D2:D4)"
    assert data["mergeData"] == [{"startRow": 6, "startColumn": 0, "endRow": 6, "endColumn": 3}]
    # SheetJS logs a warning for ODS number formats it cannot map exactly; nothing else may fail
    assert [e for e in page.errors if not e.startswith("ODS number format may be incorrect")] == []


# ---- export ----------------------------------------------------------------------------------------

def test_xlsx_round_trip(page, samples):
    open_page(page)
    import_file(page, samples["xlsx"])
    page.evaluate("S6.univerAPI.getActiveWorkbook().getActiveSheet().getRange('B2').setValue(10)")
    page.wait_for_timeout(1500)
    out = export(page, "xlsx")
    formulas = openpyxl.load_workbook(io.BytesIO(out))
    results = openpyxl.load_workbook(io.BytesIO(out), data_only=True)
    ws, rs = formulas["Data"], results["Data"]
    assert [ws[f"D{r}"].value for r in range(2, 6)] == ["=B2*C2", "=B3*C3", "=B4*C4", "=SUM(D2:D4)"]
    assert [rs[f"D{r}"].value for r in range(2, 6)] == [15, 9, 50, 74]  # recalculated after the edit
    assert [ws[f"A{r}"].value for r in range(1, 5)] == ["Item", "Äpfel", "Birnen", "Kirschen 🍒"]
    assert ws["E2"].value == datetime.datetime(2026, 1, 15) and ws["E2"].number_format == "yyyy-mm-dd"
    assert [ws[f"F{r}"].value for r in range(2, 5)] == [True, False, True]
    a1 = ws["A1"]
    assert (a1.font.b, a1.font.color.rgb, a1.fill.fgColor.rgb) == (True, "FF0000FF", "FFFFFF00")
    assert (a1.border.bottom.style, a1.alignment.horizontal, a1.alignment.vertical) == ("thin", "center", "top")
    assert (ws["A5"].font.i, ws["A5"].font.color.rgb) == (True, "FFFF0000")
    assert (ws["D5"].border.top.style, ws["D5"].font.u, ws["D5"].font.strike) == ("double", "single", True)
    assert ws["C2"].number_format == "#,##0.00"
    assert (ws["A7"].font.name, ws["A7"].font.sz, ws["A7"].alignment.wrap_text) == ("Arial", 16, True)
    assert [str(r) for r in ws.merged_cells.ranges] == ["A7:D7"]
    assert ws.column_dimensions["A"].width == pytest.approx(20, abs=0.2)
    assert ws.column_dimensions["B"].width == pytest.approx(8, abs=0.2)
    assert ws.row_dimensions[1].height == pytest.approx(30, abs=0.5)
    assert ws.freeze_panes == "A2"
    second = formulas["Second"]
    assert second["A1"].value == "=Data!D5*2" and second.sheet_properties.tabColor.rgb == "FF00B050"
    assert results["Second"]["A1"].value == 148
    assert page.errors == []


@needs_soffice
def test_ods_export(page, samples, tmp_path):
    open_page(page)
    import_file(page, samples["xlsx"])
    ods = tmp_path / "export.ods"
    ods.write_bytes(export(page, "ods"))
    # LibreOffice must be able to open it: convert back to XLSX and check values and formulas
    back = office_samples.convert(ods, "xlsx:Calc MS Excel 2007 XML", tmp_path / "back")
    ws = openpyxl.load_workbook(back)["Data"]
    assert [ws[f"A{r}"].value for r in range(1, 5)] == ["Item", "Äpfel", "Birnen", "Kirschen 🍒"]
    assert ws["D3"].value == "=B3*C3" and ws["D5"].value == "=SUM(D2:D4)"
    assert openpyxl.load_workbook(back, data_only=True)["Data"]["D5"].value == 63.5


def test_csv_export(page, samples):
    open_page(page)
    import_file(page, samples["xlsx"])
    data = export(page, "csv")
    assert data.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM, so that Excel detects the encoding
    lines = data.decode("utf-8-sig").splitlines()
    assert lines[0] == "Item,Qty,Price,Total,Date,Ok"
    assert lines[1].startswith("Äpfel,3,1.50,4.5,")


# ---- UI and helpers --------------------------------------------------------------------------------

def test_ui_open_and_save(page, samples, tmp_path):
    open_page(page)
    page.set_input_files("#open", str(samples["xlsx"]))
    page.wait_for_selector("#status:text-is('sample.xlsx')", timeout=30000)  # not "opening sample.xlsx…"
    with page.expect_download() as download:
        page.click("#export-xlsx")
    assert download.value.suggested_filename == "sample.xlsx"
    path = tmp_path / "saved.xlsx"
    download.value.save_as(path)
    assert openpyxl.load_workbook(path)["Data"]["A2"].value == "Äpfel"


@pytest.mark.parametrize("formula, rows, columns, expected", [
    ("B2*C2", 1, 0, "B3*C3"),
    ("$A$1+A1", 2, 1, "$A$1+B3"),
    ("$A1+A$1", 1, 1, "$A2+B$1"),
    ("SUM(A1:B2)", 0, 2, "SUM(C1:D2)"),
    ('"A1"&A1', 1, 0, '"A1"&A2'),
    ("Data!A1*'My Sheet'!B2", 1, 0, "Data!A2*'My Sheet'!B3"),
    ("LOG10(A1)", 1, 0, "LOG10(A2)"),
    ("Z1+1", 0, 1, "AA1+1"),
])
def test_shift_formula(page, formula, rows, columns, expected):
    open_page(page)
    assert page.evaluate("a => Suite.univerXlsx.shiftFormula(...a)", [formula, rows, columns]) == expected
