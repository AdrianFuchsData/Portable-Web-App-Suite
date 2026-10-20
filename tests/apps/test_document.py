"""Document editor (apps/document) on Univer Docs. DOCX fidelity is tested in tests/prototypes/test_s8_docx_reader.py."""
import io

from docx import Document
from playwright.sync_api import expect

from helpers import office_samples

NO_PICKERS = "window.showOpenFilePicker = undefined; window.showSaveFilePicker = undefined;"


def ready(app_page, **kw):
    page = app_page("document", init_script=NO_PICKERS, **kw)
    page.wait_for_selector("[data-testid=document] canvas", timeout=60000)
    page.wait_for_timeout(1600)
    return page


def open_file(page, path, shown=None):
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(path))
    expect(page.locator("[data-testid=file-name]")).to_have_text(shown or path.name, timeout=30000)


def text(page):
    return page.evaluate("DOC.toText(DOC.doc.save())")


def click_text(page):
    """Click into the first lines of the document (inside the canvas also on a phone, where it is 390 px wide)."""
    canvas = page.locator("[data-testid=document] canvas").first
    canvas.click(position={"x": min(400, canvas.bounding_box()["width"] / 2), "y": 120})


def test_open_edit_and_save_docx(app_page, tmp_path):
    source = office_samples.make_docx(tmp_path / "report.docx")
    page = ready(app_page)
    assert page.title() == "Document"
    open_file(page, source)
    assert text(page).startswith("Projektbericht\nNormal text with bold, italic, underlined and red.\n")
    click_text(page)
    page.keyboard.press("Control+Home")
    page.keyboard.type("Neu: ")
    expect(page.locator("[data-testid=file-name]")).to_contain_text("•")
    with page.expect_download() as d:
        page.click("[data-testid=save]")
    assert d.value.suggested_filename == "report.docx"
    saved = Document(io.BytesIO(open(d.value.path(), "rb").read()))
    paragraphs = [p.text for p in saved.paragraphs]
    assert paragraphs[0] == "Neu: Projektbericht"
    assert any(r.bold and r.text == "bold" for p in saved.paragraphs for r in p.runs)
    expect(page.locator("[data-testid=file-name]")).not_to_contain_text("•")
    assert page.errors == [] and page.requests == []


def test_typing_in_new_document_and_text_export(app_page):
    page = ready(app_page)
    click_text(page)
    page.keyboard.type("First line")
    page.keyboard.press("Enter")
    page.keyboard.type("Second line")  # (Firefox duplicates typed umlauts in Univer Docs under Playwright)
    page.click("[data-testid=save-as]")
    page.fill("[data-testid=save-as-name]", "notes")
    page.select_option("[data-testid=save-as-format]", "txt")
    with page.expect_download() as d:
        page.click("[data-testid=save-as-ok]")
    assert d.value.suggested_filename == "notes.txt"
    assert open(d.value.path(), encoding="utf-8").read() == "First line\nSecond line\n"
    expect(page.locator("[data-testid=toast]").last).to_contain_text("only the text")


def test_open_text_file(app_page, tmp_path):
    path = tmp_path / "plain.txt"
    path.write_text("Zeile 1 äöü\r\nZeile 2\r\n\r\nZeile 4", encoding="utf-8")
    page = ready(app_page)
    open_file(page, path)
    assert text(page) == "Zeile 1 äöü\nZeile 2\n\nZeile 4\n"


def test_rich_docx_and_template(app_page, tmp_path):
    rich = office_samples.make_docx_rich(tmp_path / "rich.docx")
    page = ready(app_page)
    open_file(page, rich)
    data = page.evaluate("DOC.doc.save()")
    assert "Dokumenttitel" in data["body"]["dataStream"] and data["headers"] and data["footers"]
    template = tmp_path / "vorlage.dotx"
    template.write_bytes(rich.read_bytes())
    open_file(page, template, shown="vorlage.docx")  # templates are saved as new documents


def test_language_and_theme(app_page):
    page = ready(app_page)
    page.click("[data-testid=language]")
    expect(page.locator("[data-testid=app-title]")).to_have_text("Textverarbeitung")
    expect(page.get_by_text("Einfügen", exact=True).first).to_be_visible()  # Univer's ribbon follows the language
    page.click("[data-testid=theme]")
    assert page.errors == []


def test_mobile_layout(app_page):
    page = ready(app_page, mobile=True)
    expect(page.locator("[data-testid=document] canvas").first).to_be_visible()
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
    assert page.errors == []


ODT_NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
          'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
          'xmlns:xlink="http://www.w3.org/1999/xlink" xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" office:version="1.3"')


def make_odt(path):
    """A Writer document written by hand: styles with a parent chain, heading, list, table with spans, image, link, footnote, header."""
    import zipfile
    from PIL import Image
    png = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 0, 0)).save(png, "PNG")
    content = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document-content {ODT_NS}><office:automatic-styles>
<style:style style:name="P1" style:family="paragraph" style:parent-style-name="Standard"><style:paragraph-properties fo:text-align="center"/></style:style>
<style:style style:name="T1" style:family="text"><style:text-properties fo:font-weight="bold"/></style:style>
<style:style style:name="T2" style:family="text"><style:text-properties fo:font-style="italic" fo:color="#ff0000" fo:font-size="16pt"/></style:style>
<style:style style:name="Cell" style:family="table-cell"><style:table-cell-properties fo:background-color="#ffff00"/></style:style>
<style:style style:name="Col" style:family="table-column"><style:table-column-properties style:column-width="5cm"/></style:style>
<text:list-style style:name="Num"><text:list-level-style-number text:level="1" style:num-format="1"/><text:list-level-style-number text:level="2" style:num-format="a"/></text:list-style>
</office:automatic-styles><office:body><office:text>
<text:h text:style-name="Heading_20_1" text:outline-level="1">Bericht</text:h>
<text:p text:style-name="P1">Normal <text:span text:style-name="T1">fett</text:span> und <text:span text:style-name="T2">rot</text:span><text:s text:c="2"/>Ende<text:note text:note-class="footnote"><text:note-citation>1</text:note-citation><text:note-body><text:p>Die Fußnote.</text:p></text:note-body></text:note></text:p>
<text:p text:style-name="Quote">Zitat mit <text:a xlink:type="simple" xlink:href="https://example.org/">Link</text:a></text:p>
<text:list text:style-name="Num"><text:list-item><text:p>Eins</text:p><text:list><text:list-item><text:p>Unterpunkt</text:p></text:list-item></text:list></text:list-item><text:list-item><text:p>Zwei</text:p></text:list-item></text:list>
<table:table table:name="T"><table:table-column table:style-name="Col" table:number-columns-repeated="2"/>
<table:table-row><table:table-cell table:number-columns-spanned="2" table:style-name="Cell"><text:p>Kopf</text:p></table:table-cell><table:covered-table-cell/></table:table-row>
<table:table-row><table:table-cell><text:p>A</text:p></table:table-cell><table:table-cell><text:p>B</text:p></table:table-cell></table:table-row></table:table>
<text:p><draw:frame draw:name="Bild" text:anchor-type="as-char" svg:width="2cm" svg:height="1cm"><draw:image xlink:href="Pictures/a.png"/></draw:frame></text:p>
</office:text></office:body></office:document-content>"""
    styles = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles {ODT_NS}><office:styles>
<style:default-style style:family="paragraph"><style:text-properties fo:font-size="12pt"/></style:default-style>
<style:style style:name="Standard" style:family="paragraph"/>
<style:style style:name="Quote" style:family="paragraph" style:parent-style-name="Standard"><style:paragraph-properties fo:margin-left="1cm"/><style:text-properties fo:font-style="italic"/></style:style>
<style:style style:name="Heading_20_1" style:display-name="Heading 1" style:family="paragraph" style:default-outline-level="1"><style:text-properties fo:font-size="20pt" fo:font-weight="bold"/></style:style>
</office:styles><office:automatic-styles><style:page-layout style:name="pm1"><style:page-layout-properties fo:page-width="21cm" fo:page-height="29.7cm" fo:margin-top="2cm" fo:margin-bottom="2cm" fo:margin-left="2.5cm" fo:margin-right="2.5cm"/></style:page-layout></office:automatic-styles>
<office:master-styles><style:master-page style:name="Standard" style:page-layout-name="pm1"><style:header><text:p>Kopfzeile</text:p></style:header></style:master-page></office:master-styles></office:document-styles>"""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content)
        z.writestr("styles.xml", styles)
        z.writestr("Pictures/a.png", png.getvalue())
        z.writestr("META-INF/manifest.xml", '<?xml version="1.0"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/></manifest:manifest>')
    return path


def test_open_and_save_odt(app_page, tmp_path):
    import xml.etree.ElementTree as ET
    import zipfile
    page = ready(app_page)
    open_file(page, make_odt(tmp_path / "bericht.odt"))
    content = text(page)
    assert content.startswith("Bericht\nNormal fett und rot  Ende1\nZitat mit Link\nEins\nUnterpunkt\nZwei\n")
    assert "Kopf" in content and "Die Fußnote." in content
    data = page.evaluate("DOC.doc.save()")
    body = data["body"]
    stream = body["dataStream"]
    runs = body["textRuns"]
    style_of = lambda word: next(r["ts"] for r in runs if stream[r["st"]:r["ed"]] == word)
    assert style_of("fett").get("bl") == 1 and style_of("rot").get("it") == 1 and style_of("rot")["cl"]["rgb"] == "#FF0000" and style_of("rot")["fs"] == 16
    paragraphs = body["paragraphs"]
    assert paragraphs[0]["paragraphStyle"]["namedStyleType"] == 4  # heading 1
    assert paragraphs[1]["paragraphStyle"]["horizontalAlign"] == 2
    lists = [p["bullet"] for p in paragraphs if p.get("bullet")]
    assert [(b["listType"], b["nestingLevel"]) for b in lists] == [("ORDER_LIST", 0), ("ORDER_LIST", 1), ("ORDER_LIST", 0)]
    assert any(r["properties"]["url"] == "https://example.org/" for r in body["customRanges"])
    [table] = data["tableSource"].values()
    assert table["tableRows"][0]["tableCells"][0]["columnSpan"] == 2 and table["tableRows"][0]["tableCells"][0]["backgroundColor"]["rgb"] == "#FFFF00"
    assert round(table["tableColumns"][0]["size"]["width"]["v"]) == 189
    assert len(data["drawings"]) == 1 and round(data["documentStyle"]["marginLeft"]) == 94
    assert any("Kopfzeile" in h["body"]["dataStream"] for h in data["headers"].values())

    # saved as ODT again
    with page.expect_download() as d:
        page.click("[data-testid=save]")
    assert d.value.suggested_filename == "bericht.odt"
    with zipfile.ZipFile(d.value.path()) as z:
        assert z.namelist()[0] == "mimetype" and z.read("mimetype") == b"application/vnd.oasis.opendocument.text"
        assert z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
        root = ET.fromstring(z.read("content.xml"))
        styles = ET.fromstring(z.read("styles.xml"))
        assert any(n.startswith("Pictures/") for n in z.namelist())
    ns = {"text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0", "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
          "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0", "fo": "urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0",
          "xlink": "http://www.w3.org/1999/xlink"}
    heading = root.find(".//text:h", ns)
    assert heading.get("{%s}outline-level" % ns["text"]) == "1" and "".join(heading.itertext()) == "Bericht"
    spans = {"".join(s.itertext()): s.get("{%s}style-name" % ns["text"]) for s in root.iter("{%s}span" % ns["text"])}
    auto = {s.get("{%s}name" % ns["style"]): s for s in root.iter("{%s}style" % ns["style"])}
    bold = auto[spans["fett"]].find("style:text-properties", ns)
    assert bold.get("{%s}font-weight" % ns["fo"]) == "bold"
    assert root.find(".//text:a", ns).get("{%s}href" % ns["xlink"]) == "https://example.org/"
    items = ["".join(p.itertext()) for p in root.find(".//text:list", ns).iter("{%s}p" % ns["text"])]
    assert items == ["Eins", "Unterpunkt", "Zwei"]
    cells = root.findall(".//table:table-row", ns)[0]
    assert cells[0].get("{%s}number-columns-spanned" % ns["table"]) == "2" and cells[1].tag.endswith("covered-table-cell")
    assert "Kopfzeile" in "".join(styles.find(".//style:header", ns).itertext())
    # and read back
    path = tmp_path / "wieder.odt"
    d.value.save_as(path)
    open_file(page, path)
    assert text(page).startswith("Bericht\nNormal fett und rot  Ende1\n")
    assert page.errors == []


def make_docx_with_headings_and_footnote(path):
    """python-docx has no footnotes: the footnote part is added to the package by hand."""
    import re
    import zipfile
    d = Document()
    d.add_heading("Einleitung", 1)
    p = d.add_paragraph("Ein Satz mit Fußnote")
    d.add_heading("Details", 2)
    d.add_paragraph("Mehr Text.")
    d.add_heading("Schluss", 1)
    d.save(path)
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    doc = files["word/document.xml"].decode()
    doc = doc.replace("<w:t>Ein Satz mit Fußnote</w:t></w:r>", '<w:t>Ein Satz mit Fußnote</w:t></w:r><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="2"/></w:r>', 1)
    files["word/document.xml"] = doc.encode()
    files["word/footnotes.xml"] = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{w}">'
                                   '<w:footnote w:type="separator" w:id="0"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                                   '<w:footnote w:id="2"><w:p><w:r><w:t>Die Quelle der Fußnote.</w:t></w:r></w:p></w:footnote></w:footnotes>').encode()
    rels = files["word/_rels/document.xml.rels"].decode()
    files["word/_rels/document.xml.rels"] = rels.replace("</Relationships>", '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>').encode()
    types = files["[Content_Types].xml"].decode()
    files["[Content_Types].xml"] = types.replace("</Types>", '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>').encode()
    with zipfile.ZipFile(path, "w") as z:
        for n, data in files.items():
            z.writestr(n, data)
    return path


def test_table_of_contents_and_footnotes(app_page, tmp_path):
    import zipfile
    page = ready(app_page)
    open_file(page, make_docx_with_headings_and_footnote(tmp_path / "kapitel.docx"))
    content = text(page)
    assert "Ein Satz mit Fußnote1\n" in content and content.rstrip().endswith("—————\n1 Die Quelle der Fußnote.")
    page.click("[data-testid=toc]")
    expect(page.locator("[data-testid=toast]").last).to_contain_text("3 headings inserted")
    assert text(page).startswith("Contents\nEinleitung\nDetails\nSchluss\nEinleitung\n")
    indents = page.evaluate("DOC.doc.save().body.paragraphs.slice(1, 4).map((p) => (p.paragraphStyle.indentStart || { v: 0 }).v)")
    assert indents == [0, 18, 0]
    # a second click updates it in place
    page.click("[data-testid=toc]")
    expect(page.locator("[data-testid=toast]").last).to_contain_text("updated (3 headings)")
    assert text(page).count("Contents\n") == 1
    # DOCX: internal links to bookmarks on the headings
    with page.expect_download() as d:
        page.click("[data-testid=save]")
    with zipfile.ZipFile(d.value.path()) as z:
        xml = z.read("word/document.xml").decode()
    assert 'w:anchor="toc-1"' in xml and 'w:anchor="toc-2"' in xml and 'w:name="toc-1"' in xml and 'w:name="toc-3"' in xml
    # ODT: the same with text:bookmark
    page.click("[data-testid=save-as]")
    page.select_option("[data-testid=save-as-format]", "odt")
    with page.expect_download() as d:
        page.click("[data-testid=save-as-ok]")
    with zipfile.ZipFile(d.value.path()) as z:
        content_xml = z.read("content.xml").decode()
    import re
    assert re.search(r'<text:bookmark text:name="toc-2"/>(<text:span[^>]*>)?Details', content_xml) and 'xlink:href="#toc-2"' in content_xml
    assert page.errors == []


def test_print_and_html_export(app_page, tmp_path):
    page = ready(app_page)
    page.evaluate("window.printed = 0")
    open_file(page, make_odt(tmp_path / "bericht.odt"))
    page.click("[data-testid=toc]")
    # printing renders the document into a hidden frame with the page size and margins of the document
    # (Firefox under Playwright blocks the page in print(): there the frame is caught before it loads)
    page.evaluate("""(stop) => { new MutationObserver((list) => { for (const m of list) for (const n of m.addedNodes)
      if (n.tagName === 'IFRAME') { window.printed += 1; window.printedHtml = n.srcdoc; if (stop) n.onload = null; } })
      .observe(document.body, { childList: true }); }""", page.context.browser.browser_type.name == "firefox")
    page.click("[data-testid=print]")
    page.wait_for_function("() => window.printed === 1")
    printed = page.evaluate("window.printedHtml")
    assert "@page { size: " in printed and "margin: " in printed
    assert '<h1 id="toc-1"' in printed and '<a href="#toc-1">Bericht</a>' in printed
    assert 'font-weight:bold;font-size:12pt">fett</span>' in printed and '<a href="https://example.org/">' in printed
    assert "<ol><li>" in printed and 'colspan="2"' in printed and "<img src=\"data:image/" in printed
    assert 'class="doc-header"' in printed and "Kopfzeile" in printed
    # saved as a web page
    page.click("[data-testid=save-as]")
    page.select_option("[data-testid=save-as-format]", "html")
    expect(page.locator("[data-testid=save-as-dialog], dialog[open]").first).to_contain_text("not the page layout")
    with page.expect_download() as d:
        page.click("[data-testid=save-as-ok]")
    assert d.value.suggested_filename == "bericht.html"
    saved = open(d.value.path(), encoding="utf-8").read()
    assert saved.startswith("<!doctype html>") and "<title>bericht</title>" in saved and "Zitat mit " in saved
    assert page.errors == []


def test_spelling(app_page, tmp_path):
    import zipfile
    page = ready(app_page)
    page.evaluate("DOC.load(DOC.fromText(\"Hello wrold and more teext here.\\nThe wrold is big, NASA and e.g. https://exmaple.org don't care.\"), 'spell.docx')")
    page.wait_for_timeout(500)
    page.click("[data-testid=spell]")
    expect(page.locator("[data-testid=spell-status]")).to_have_text("2 unknown words", timeout=20000)
    expect(page.locator("[data-testid=spell-word]")).to_have_text(["wrold ×2", "teext"])
    # a word selects its occurrence in the document and shows suggestions; a suggestion replaces all occurrences
    page.click("[data-testid=spell-word] >> nth=0")
    expect(page.locator("[data-testid=spell-suggestion]", has_text="world")).to_be_visible()
    assert page.evaluate("DOC.univerAPI._injector.get(DOC.selection).getActiveTextRange().startOffset") == 6
    page.click("[data-testid=spell-suggestion] >> text=world")
    expect(page.locator("[data-testid=spell-status]")).to_have_text("1 unknown word")
    assert text(page).startswith("Hello world and more teext here.\nThe world is big")
    expect(page.locator("[data-testid=file-name]")).to_contain_text("•")
    # added to the personal dictionary (kept for the next session)
    page.click("[data-testid=spell-word]")
    page.click("[data-testid=spell-add]")
    expect(page.locator("[data-testid=spell-status]")).to_have_text("No spelling errors found.")
    assert page.evaluate("JSON.parse(localStorage.getItem('suite.spell.words'))") == ["teext"]

    # a dictionary from a LibreOffice extension (ISO-8859-1), kept in IndexedDB
    aff = "SET ISO8859-1\nTRY esianrtolcdugmphbyfvkwzESIANRTOLCDUGMPHBYFVKWZ\n".encode("latin-1")
    dic = "3\nHallo\nWelt\nGrüße\n".encode("latin-1")
    oxt = tmp_path / "dict-de.oxt"
    with zipfile.ZipFile(oxt, "w") as z:
        z.writestr("META-INF/manifest.xml", "<manifest/>")
        z.writestr("dictionaries/de_TEST.aff", aff)
        z.writestr("dictionaries/de_TEST.dic", dic)
        z.writestr("dictionaries/hyph_de_TEST.dic", b"ISO8859-1\n")
    page.evaluate("DOC.load(DOC.fromText('Hallo Welt, viele Grüße und Grüsse'), 'de.docx')")
    page.click("[data-testid=spell-recheck]")
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=spell-load]")
    chooser.value.set_files(str(oxt))
    expect(page.locator("[data-testid=spell-language]")).to_have_value("de_TEST")
    expect(page.locator("[data-testid=spell-word]")).to_have_text(["viele", "und", "Grüsse"])
    # ignored for this session
    page.click("[data-testid=spell-word] >> text=und")
    page.click("[data-testid=spell-ignore]")
    expect(page.locator("[data-testid=spell-word]")).to_have_text(["viele", "Grüsse"])
    # still there after a reload; .aff and .dic files work as well
    page.reload()
    page.wait_for_selector("[data-testid=document] canvas", timeout=60000)
    page.keyboard.press("F7")
    expect(page.locator("[data-testid=spell-language] option")).to_have_count(3)  # English, German, de_TEST
    expect(page.locator("[data-testid=spell-language]")).to_have_value("de_TEST")
    page.click("[data-testid=spell-remove]")
    expect(page.locator("[data-testid=spell-language] option")).to_have_count(2)
    # built-in German: compounds and umlauts are known (Hunspell)
    page.evaluate("DOC.load(DOC.fromText('Die Rechtschreibprüfung der Kindergartenplätze im Gesundheitsministerium ist schön, aber Fehlerr bleiben.'), 'de2.docx')")
    page.select_option("[data-testid=spell-language]", "de_DE")
    page.click("[data-testid=spell-recheck]")
    expect(page.locator("[data-testid=spell-word]")).to_have_text(["Fehlerr"], timeout=20000)
    page.click("[data-testid=spell-word]")
    expect(page.locator("[data-testid=spell-suggestion]").first).to_have_text("Fehler")
    (tmp_path / "en_X.aff").write_text("SET UTF-8\n")
    (tmp_path / "en_X.dic").write_text("1\nzyx\n")
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=spell-load]")
    chooser.value.set_files([str(tmp_path / "en_X.aff"), str(tmp_path / "en_X.dic")])
    expect(page.locator("[data-testid=spell-language]")).to_have_value("en_X")
    page.click("[data-testid=spell-close]")
    expect(page.locator("[data-testid=spell-panel]")).to_have_count(0)
    assert page.errors == []
