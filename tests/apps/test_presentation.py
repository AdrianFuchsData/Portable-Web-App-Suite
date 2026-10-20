"""Presentation editor (apps/presentation). PPTX import/export fidelity is tested in tests/prototypes/test_s7_slides.py."""
import io

from playwright.sync_api import expect
from pptx import Presentation
from pptx.util import Emu

from helpers import office_samples

NO_PICKERS = "window.showOpenFilePicker = undefined; window.showSaveFilePicker = undefined;"
PNG = bytes.fromhex("89504e470d0a1a0a0000000d4948445200000001000000010806000000"
                    "1f15c4890000000d49444154789c6360f8cf000003010100c9fe92ef0000000049454e44ae426082")


def sample(tmp_path):
    path = tmp_path / "deck.pptx"
    office_samples.make_pptx(path)
    return path


def ready(app_page, **kw):
    page = app_page("presentation", init_script=NO_PICKERS, **kw)
    expect(page.locator("[data-testid=thumb]")).to_have_count(1)
    return page


def open_deck(page, path):
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(path))
    expect(page.locator("[data-testid=file-name]")).to_have_text(path.name)


def save(page, tmp_path):
    with page.expect_download() as d:
        page.click("[data-testid=save]")
    return Presentation(io.BytesIO(open(d.value.path(), "rb").read())), d.value.suggested_filename


def elements(page):
    return page.evaluate("PR.state.deck.slides[PR.state.current].elements.map(e => ({ id: e.id, type: e.type, x: e.x, y: e.y, w: e.w, h: e.h }))")


def stage_el(page, i):
    return page.locator("[data-testid=stage] .el").nth(i)


def test_open_edit_and_save(app_page, tmp_path):
    page = ready(app_page)
    assert page.title() == "Presentation"
    open_deck(page, sample(tmp_path))
    expect(page.locator("[data-testid=thumb]")).to_have_count(3)
    stage_el(page, 1).click()
    expect(page.locator("[data-testid=panel]")).to_contain_text("Subtitle")
    page.fill("[data-testid=el-x]", "100")
    page.locator("[data-testid=el-x]").blur()
    page.fill("[data-testid=font-size]", "24")
    page.locator("[data-testid=font-size]").blur()
    page.click("[data-testid=fmt-bold]")
    expect(page.locator("[data-testid=file-name]")).to_contain_text("•")
    prs, name = save(page, tmp_path)
    assert name == "deck.pptx"
    sub = prs.slides[0].shapes[1]
    assert round(Emu(sub.left).inches * 96) == 100
    runs = [r for p in sub.text_frame.paragraphs for r in p.runs]
    assert runs[0].text == "Vertraulich – äöü" and runs[0].font.bold and runs[0].font.size.pt == 24
    expect(page.locator("[data-testid=file-name]")).not_to_contain_text("•")
    assert page.errors == [] and page.requests == []


def test_add_elements_and_edit_text(app_page, tmp_path):
    page = ready(app_page)
    page.click("[data-testid=add-text]")
    expect(page.locator("[data-testid=text-editor]")).to_be_focused()
    page.keyboard.press("Control+a")
    page.keyboard.type("Hello slides")
    page.locator("[data-testid=stage]").click(position={"x": 5, "y": 5})  # leave the editor
    page.locator("[data-testid=shape-menu] summary").click()
    page.click("[data-testid=add-ellipse]")
    page.fill("[data-testid=el-fill]", "#ff0000")
    page.click("[data-testid=add-table]")
    table = stage_el(page, 3)
    table.locator("td").first.dblclick()
    expect(page.locator("[data-testid=cell-editor]")).to_be_focused()
    page.keyboard.type("Cell A1")
    page.locator("[data-testid=stage]").click(position={"x": 5, "y": 5})
    image = tmp_path / "dot.png"
    image.write_bytes(PNG)
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=add-image]")
    chooser.value.set_files(str(image))
    expect(page.locator("[data-testid=stage] .el")).to_have_count(5)
    prs, _ = save(page, tmp_path)
    shapes = prs.slides[0].shapes
    texts = [s.text_frame.text for s in shapes if s.has_text_frame]
    assert "Hello slides" in texts
    assert any(s.has_table and s.table.cell(0, 0).text == "Cell A1" for s in shapes)
    assert any(s.shape_type == 13 for s in shapes)  # picture
    assert any(s.has_text_frame is not None and getattr(s, "fill", None) is not None and s.fill.type == 1 and str(s.fill.fore_color.rgb) == "FF0000" for s in shapes if s.shape_type == 1)


def test_slides_undo_and_keyboard(app_page):
    page = ready(app_page)
    page.click("[data-testid=add-slide]")
    expect(page.locator("[data-testid=thumb]")).to_have_count(2)
    page.click("[data-testid=duplicate-slide]")
    expect(page.locator("[data-testid=thumb]")).to_have_count(3)
    page.click("[data-testid=slide-up]")
    expect(page.locator("[data-testid=thumb].active")).to_have_count(1)
    assert page.evaluate("PR.state.current") == 1
    page.click("[data-testid=delete-slide]")
    expect(page.locator("[data-testid=thumb]")).to_have_count(2)
    page.click("[data-testid=undo]")
    expect(page.locator("[data-testid=thumb]")).to_have_count(3)
    page.click("[data-testid=redo]")
    expect(page.locator("[data-testid=thumb]")).to_have_count(2)

    page.locator("[data-testid=thumb]").first.click()
    stage_el(page, 0).click()
    before = elements(page)[0]
    page.keyboard.press("ArrowRight")
    page.keyboard.press("Shift+ArrowDown")
    after = elements(page)[0]
    assert (after["x"] - before["x"], after["y"] - before["y"]) == (1, 10)
    page.keyboard.press("Control+c")
    page.keyboard.press("Control+v")
    expect(page.locator("[data-testid=stage] .el")).to_have_count(2)
    page.keyboard.press("Delete")
    expect(page.locator("[data-testid=stage] .el")).to_have_count(1)


def test_move_with_mouse(app_page, tmp_path):
    page = ready(app_page)
    stage_el(page, 0).click()
    box = stage_el(page, 0).bounding_box()
    before = elements(page)[0]
    page.mouse.move(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)
    page.mouse.down()
    page.mouse.move(box["x"] + box["width"] / 2 + 60, box["y"] + box["height"] / 2 + 40, steps=8)
    page.mouse.up()
    after = elements(page)[0]
    assert after["x"] > before["x"] + 20 and after["y"] > before["y"] + 10
    expect(page.locator("[data-testid=file-name]")).to_contain_text("•")


def test_present(app_page, tmp_path):
    page = ready(app_page)
    open_deck(page, sample(tmp_path))
    page.click("[data-testid=present]")
    presenter = page.locator("[data-testid=presenter]")
    expect(presenter).to_be_visible()
    expect(presenter).to_have_attribute("data-index", "0")
    page.keyboard.press("ArrowRight")
    expect(presenter).to_have_attribute("data-index", "1")
    page.keyboard.press("End")
    expect(presenter).to_have_attribute("data-index", "2")
    page.keyboard.press("ArrowRight")  # past the end: the show ends
    expect(presenter).to_be_hidden()
    page.keyboard.press("F5")
    expect(presenter).to_have_attribute("data-index", "0")
    page.keyboard.press("Escape")
    expect(presenter).to_be_hidden()


def test_save_as_and_template(app_page, tmp_path):
    page = ready(app_page)
    template = tmp_path / "vorlage.potx"
    office_samples.make_pptx(template)
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(template))
    expect(page.locator("[data-testid=file-name]")).to_have_text("vorlage.pptx")
    page.click("[data-testid=save-as]")
    page.fill("[data-testid=save-as-name]", "bericht")
    with page.expect_download() as d:
        page.click("[data-testid=save-as-ok]")
    assert d.value.suggested_filename == "bericht.pptx"
    expect(page.locator("[data-testid=file-name]")).to_have_text("bericht.pptx")


def test_german_interface(app_page):
    page = app_page("presentation", init_script="localStorage.setItem('suite.lang', 'de');")
    expect(page.locator("[data-testid=present]")).to_contain_text("Vorführen")


def test_mobile_layout(app_page):
    page = ready(app_page, mobile=True)
    expect(page.locator("[data-testid=panel]")).to_be_hidden()
    page.click("[data-testid=toggle-panels]")
    expect(page.locator("[data-testid=panel]")).to_be_visible()
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
    assert page.errors == []


def test_notes_and_presenter_view(app_page, tmp_path):
    path = sample(tmp_path)
    prs = Presentation(path)
    prs.slides[1].notes_slide.notes_text_frame.text = "Imported note\nsecond line"
    prs.save(path)
    page = ready(app_page)
    open_deck(page, path)
    assert page.evaluate("PR.state.deck.slides[1].notes") == "Imported note\nsecond line"
    page.fill("[data-testid=slide-notes]", "Say hello")
    assert page.evaluate("PR.state.deck.slides[0].notes") == "Say hello"
    # presenter view: console here, slides in the audience window
    with page.context.expect_page() as popup:
        page.click("[data-testid=present-console]")
    audience = popup.value
    console = page.locator("[data-testid=console]")
    expect(console).to_be_visible()
    expect(page.locator("[data-testid=console-count]")).to_have_text("Slide 1 of 3")
    expect(page.locator("[data-testid=console-notes]")).to_have_text("Say hello")
    expect(page.locator("[data-testid=console-next] .slide")).to_have_count(1)
    expect(audience.locator("#stage .slide")).to_have_count(1)
    page.keyboard.press("ArrowRight")
    expect(page.locator("[data-testid=console-count]")).to_have_text("Slide 2 of 3")
    expect(page.locator("[data-testid=console-notes]")).to_have_text("Imported note\nsecond line")
    # keys in the audience window work as well
    audience.keyboard.press("ArrowRight")
    expect(page.locator("[data-testid=console-count]")).to_have_text("Slide 3 of 3")
    expect(page.locator("[data-testid=console-next]")).to_contain_text("End of the presentation")
    page.keyboard.press("ArrowRight")  # stays on the last slide
    expect(page.locator("[data-testid=console-count]")).to_have_text("Slide 3 of 3")
    expect(page.locator("[data-testid=console-timer]")).to_have_text("00:00")
    page.click("[data-testid=console-end]")
    expect(console).to_be_hidden()
    assert audience.is_closed()
    # notes are saved in the PPTX file
    saved, _ = save(page, tmp_path)
    assert saved.slides[0].notes_slide.notes_text_frame.text == "Say hello"
    assert saved.slides[1].notes_slide.notes_text_frame.text == "Imported note\nsecond line"
    assert page.errors == []


def test_transitions_and_animations(app_page, tmp_path):
    page = ready(app_page)
    open_deck(page, sample(tmp_path))
    # transition of slide 1: push from below, 0.5 s, next slide after 3 s
    page.click("[data-testid=transition-section]")
    page.select_option("[data-testid=transition-type]", "push")
    page.select_option("[data-testid=transition-dir]", "u")
    page.fill("[data-testid=transition-duration]", "0.5")
    page.press("[data-testid=transition-duration]", "Enter")
    page.fill("[data-testid=transition-advance]", "3")
    page.press("[data-testid=transition-advance]", "Enter")
    assert page.evaluate("PR.state.deck.slides[0].transition") == {"type": "push", "dir": "u", "duration": 0.5, "advanceAfter": 3}
    # entrance animations: the title flies in on click, the body fades in after it
    ids = [e["id"] for e in elements(page)]
    page.evaluate("(id) => PR.select(id)", ids[0])
    page.select_option("[data-testid=el-animation]", "fly")
    page.evaluate("(id) => PR.select(id)", ids[1])
    page.select_option("[data-testid=el-animation]", "fade")
    page.select_option("[data-testid=el-animation-start]", "after")
    assert page.evaluate("PR.state.deck.slides[0].animations") == [{"id": ids[0], "effect": "fly", "start": "click"}, {"id": ids[1], "effect": "fade", "start": "after"}]
    page.evaluate("PR.select(null)")
    expect(page.locator("[data-testid=animation-item]")).to_have_count(2)

    # saved into the PPTX: p:transition and p:timing referring to the shapes
    prs, _ = save(page, tmp_path)
    from lxml import etree
    sld = etree.fromstring(etree.tostring(prs.slides[0]._element))  # a plain lxml element (python-pptx's xpath() has no namespaces argument)
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    [tr] = sld.xpath("p:transition", namespaces=ns)
    assert (tr.get("spd"), tr.get("advTm"), tr[0].tag.split("}")[1], tr[0].get("dir")) == ("fast", "3000", "push", "u")
    effects = sld.xpath(".//p:cTn[@presetClass='entr']", namespaces=ns)
    targets = [(e.get("presetID"), e.get("nodeType"), e.xpath(".//p:spTgt/@spid", namespaces=ns)[0]) for e in effects]
    assert [(p_, n) for p_, n, _ in targets] == [("2", "clickEffect"), ("10", "afterEffect")]
    shape_ids = [str(s.shape_id) for s in prs.slides[0].shapes]
    assert targets[0][2] == shape_ids[0] and targets[1][2] == shape_ids[1]
    children = [c.tag.split("}")[1] for c in sld]
    assert children.index("transition") == children.index("clrMapOvr") + 1 and children.index("timing") == children.index("transition") + 1
    # and read again
    path = tmp_path / "motion.pptx"
    prs.save(path)
    page2 = ready(app_page)
    open_deck(page2, path)
    assert page2.evaluate("PR.state.deck.slides[0].transition") == {"type": "push", "dir": "u", "duration": 0.5, "advanceAfter": 3}
    assert page2.evaluate("PR.state.deck.slides[0].animations.map((a) => [a.effect, a.start])") == [["fly", "click"], ["fade", "after"]]

    # the show: animated elements wait for the click, then the next slide
    page.evaluate("PR.state.deck.slides[0].transition.advanceAfter = undefined; PR.state.deck.slides[1].transition = { type: 'fade', dir: 'l', duration: 0.3, advanceAfter: 1 }")
    page.click("[data-testid=present]")
    presenter = page.locator("[data-testid=presenter]")
    expect(presenter).to_have_attribute("data-step", "0")
    visibility = lambda: page.evaluate("(id) => getComputedStyle(document.querySelector(`[data-testid=presenter] .slide:last-child [data-id='${id}']`)).visibility", ids[0])
    assert visibility() == "hidden"
    page.keyboard.press("ArrowRight")
    expect(presenter).to_have_attribute("data-step", "1")
    expect(presenter).to_have_attribute("data-index", "0")
    assert visibility() == "visible"
    page.keyboard.press("ArrowRight")
    expect(presenter).to_have_attribute("data-index", "1")
    expect(presenter).to_have_attribute("data-index", "2", timeout=4000)  # slide 2 advances by itself after 1 s
    page.keyboard.press("ArrowLeft")
    page.keyboard.press("ArrowLeft")
    expect(presenter).to_have_attribute("data-index", "0")
    expect(presenter).to_have_attribute("data-step", "1")  # back: the slide is complete
    page.keyboard.press("ArrowLeft")
    expect(presenter).to_have_attribute("data-step", "0")
    page.keyboard.press("Escape")
    assert page.errors == [] and page2.errors == []


def test_powerpoint_transition_markup(app_page, tmp_path):
    from lxml import etree
    path = sample(tmp_path)
    prs = Presentation(path)
    alt = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
           '<mc:Choice xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" Requires="p14"><p:transition spd="slow" p14:dur="1250"><p:wipe dir="d"/></p:transition></mc:Choice>'
           '<mc:Fallback><p:transition spd="slow"><p:wipe dir="d"/></p:transition></mc:Fallback></mc:AlternateContent>')
    sld = prs.slides[1]._element
    sld.insert(list(sld).index(sld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}clrMapOvr")) + 1, etree.fromstring(alt))
    prs.save(path)
    page = ready(app_page)
    open_deck(page, path)
    assert page.evaluate("PR.state.deck.slides[1].transition") == {"type": "wipe", "dir": "d", "duration": 1.25}
    assert page.errors == []


def test_open_web_slides(app_page, tmp_path):
    html = """<!doctype html><html><head><title>Web</title></head><body><div class="reveal"><div class="slides">
<section data-background-color="#222222" data-transition="zoom"><h1>Willkommen</h1><p>Eine <strong>Web</strong>-Präsentation</p><aside class="notes">Begrüßung</aside></section>
<section><section><h2>Punkte</h2><ul><li class="fragment">Erstens</li><li>Zweitens</li></ul></section><section><h3>Code</h3><pre><code>print("hi")</code></pre></section></section>
<section data-markdown><textarea data-template>## Markdown
- a
- **b**</textarea></section>
</div></div></body></html>"""
    path = tmp_path / "web.html"
    path.write_text(html, encoding="utf-8")
    page = ready(app_page)
    open_deck_any = page.expect_file_chooser
    with open_deck_any() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(path))
    expect(page.locator("[data-testid=file-name]")).to_have_text("web.pptx")
    expect(page.locator("[data-testid=thumb]")).to_have_count(4)  # the vertical slide follows its stack
    deck = page.evaluate("PR.state.deck.slides.map((s) => ({ bg: s.background, notes: s.notes, tr: s.transition && s.transition.type, texts: s.elements.map((e) => e.text ? e.text.paragraphs.map((p) => p.runs.map((r) => r.text).join('')).join('/') : e.type), bullets: s.elements.flatMap((e) => e.text ? e.text.paragraphs.filter((p) => p.bullet).map((p) => p.runs.map((r) => r.text).join('')) : []), anims: (s.animations || []).length }))")
    assert deck[0]["bg"] == "#222222" and deck[0]["notes"] == "Begrüßung" and deck[0]["tr"] == "zoom" and deck[0]["texts"][0] == "Willkommen"
    assert deck[1]["bullets"] == ["Erstens", "Zweitens"] and deck[1]["anims"] == 1
    assert deck[2]["texts"] == ["Code", 'print("hi")']
    assert deck[3]["texts"][0] == "Markdown" and deck[3]["bullets"] == ["a", "b"]
    white = page.evaluate("PR.state.deck.slides[0].elements[0].text.paragraphs[0].runs[0].color")
    assert white == "#FFFFFF"  # light text on the dark background
    assert page.errors == []
