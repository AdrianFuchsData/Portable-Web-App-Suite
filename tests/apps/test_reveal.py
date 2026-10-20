"""Web slides (apps/reveal): visual reveal.js editor, saved presentations work offline, Markdown/Jupyter/reveal.js import."""
import json
import re

from playwright.sync_api import expect

from helpers.paths import dist_url

NO_PICKERS = "window.showOpenFilePicker = undefined; window.showSaveFilePicker = undefined;"


def ready(app_page, **kw):
    page = app_page("reveal", init_script=NO_PICKERS, **kw)
    wait_preview(page)
    return page


def wait_preview(page, before=None):
    """Wait until the preview deck is (re)built and reveal.js is ready."""
    for _ in range(100):
        version = page.evaluate("document.body.dataset.preview || ''")
        if version and version != before:
            break
        page.wait_for_timeout(100)
    return page.frame_locator("[data-testid=preview]")


def preview_version(page):
    return page.evaluate("document.body.dataset.preview || ''")


def save_html(page, tmp_path, name="saved.html"):
    with page.expect_download() as d:
        page.click("[data-testid=save]")
    path = tmp_path / name
    d.value.save_as(path)
    return path


def open_file(page, path):
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(path))


def test_edit_slides_and_save_presentation(app_page, tmp_path, browser):
    page = ready(app_page)
    assert page.title().endswith("Web slides")
    expect(page.locator("[data-testid=slide-item]")).to_have_count(2)
    frame = page.frame_locator("[data-testid=preview]")
    expect(frame.locator("section.present h1")).to_have_text("Presentation")

    # edit the title directly in the slide
    v = preview_version(page)
    frame.locator("section.present h1").dblclick()
    expect(page.locator("[data-testid=format-bar]")).to_be_visible()
    page.keyboard.press("Control+a")
    page.keyboard.type("My talk")
    page.click("[data-testid=fmt-done]")
    frame = wait_preview(page, v)
    expect(frame.locator("section.present h1")).to_have_text("My talk")
    expect(page.locator("[data-testid=slide-item]").first).to_contain_text("My talk")

    # a code slide with a highlighted step, background and transition
    v = preview_version(page)
    page.click("[data-testid=add-slide] summary")
    page.click("[data-testid=add-slide-code]")
    frame = wait_preview(page, v)
    expect(page.locator("[data-testid=slide-item]")).to_have_count(3)
    frame.locator("section.present pre").click()
    expect(page.locator("[data-testid=element-inspector]")).to_be_visible()
    page.select_option("[data-testid=el-language]", "python")
    page.fill("[data-testid=el-code]", "print('hi')\nx = 1")
    page.fill("[data-testid=el-lines]", "1|2")
    page.click("[data-testid=tab-slide]")
    page.wait_for_selector("[data-testid=slide-bg-color]")
    page.evaluate("() => { const el = document.querySelector('[data-testid=slide-bg-color]'); el.value = '#123456'; el.dispatchEvent(new Event('input', { bubbles: true })); }")
    page.select_option("[data-testid=slide-transition]", "zoom")
    page.fill("[data-testid=slide-notes]", "Explain the loop")
    # a list that appears point by point (the code slide was inserted after the current, first slide)
    page.click("[data-testid=slide-item] >> nth=2")
    frame = page.frame_locator("[data-testid=preview]")
    frame.locator("section.present ul").click()
    page.select_option("[data-testid=el-item-fragment]", "fade-up")
    page.click("[data-testid=tab-deck]")
    page.select_option("[data-testid=deck-theme]", "moon")
    page.check("[data-testid=deck-slide-number]")
    page.fill("[data-testid=deck-title]", "Talk")

    saved = save_html(page, tmp_path)
    text = saved.read_text(encoding="utf-8")
    assert "<title>Talk</title>" in text
    assert 'data-background-color="#123456"' in text and 'data-transition="zoom"' in text
    assert '<code class="language-python" data-trim data-line-numbers="1|2">' in text
    assert '<li class="fragment fade-up">' in text and '<aside class="notes">Explain the loop</aside>' in text
    assert '"slideNumber":true' in text and "Source Sans Pro" not in text  # moon uses League Gothic
    assert not re.search(r'<(script|link)[^>]+(src|href)="(https?:)?//', text)  # self-contained

    # the saved file is a working reveal.js presentation, offline
    context = browser.new_context(viewport={"width": 1280, "height": 800})  # slide view (phones get reveal.js's scroll view)
    deck_page = context.new_page()
    errors = []
    deck_page.on("pageerror", lambda e: errors.append(str(e)))
    deck_page.on("request", lambda r: not r.url.startswith(("file:", "data:", "blob:", "about:")) and errors.append(r.url))
    deck_page.goto(saved.as_uri())
    deck_page.wait_for_function("window.Reveal && Reveal.isReady()")
    assert deck_page.evaluate("Reveal.getTotalSlides()") == 3
    deck_page.keyboard.press("ArrowRight")
    assert deck_page.evaluate("Reveal.getIndices().h") == 1
    assert deck_page.evaluate("document.querySelectorAll('.reveal pre code.hljs').length") == 2  # highlighted, one block per step (1|2)
    assert errors == []
    context.close()

    # and it opens again with the full editor model
    page.click("[data-testid=new]")
    open_file(page, saved)
    expect(page.locator("[data-testid=slide-item]")).to_have_count(3)
    assert page.evaluate("RV.state.deck.theme") == "moon"
    assert page.evaluate("RV.state.deck.slides[1].blocks[1].lineNumbers") == "1|2"
    assert page.errors == []


def test_slide_structure_and_undo(app_page):
    page = ready(app_page)
    page.click("[data-testid=add-sub] summary")
    page.click("[data-testid=add-sub-section]")
    expect(page.locator("[data-testid=slide-item].sub")).to_have_count(1)
    expect(page.locator("[data-testid=slide-item].active .rv-num")).to_have_text("1.2")
    page.click("[data-testid=slide-outdent]")
    expect(page.locator("[data-testid=slide-item].sub")).to_have_count(0)
    expect(page.locator("[data-testid=slide-item]")).to_have_count(3)
    page.click("[data-testid=slide-indent]")
    expect(page.locator("[data-testid=slide-item].sub")).to_have_count(1)
    page.click("[data-testid=slide-duplicate]")
    expect(page.locator("[data-testid=slide-item]")).to_have_count(4)
    page.click("[data-testid=undo]")
    expect(page.locator("[data-testid=slide-item]")).to_have_count(3)
    page.click("[data-testid=undo]")
    expect(page.locator("[data-testid=slide-item].sub")).to_have_count(0)
    page.click("[data-testid=redo]")
    expect(page.locator("[data-testid=slide-item].sub")).to_have_count(1)
    page.click("[data-testid=slide-delete]")
    expect(page.locator("[data-testid=slide-item]")).to_have_count(2)
    assert page.errors == []


def test_blocks_markdown_and_html_slides(app_page, tmp_path):
    page = ready(app_page)
    page.click("[data-testid=slide-item] >> nth=1")
    for kind in ("table", "quote", "html"):
        page.click("[data-testid=add-block] summary")
        page.click(f"[data-testid=add-block-{kind}]")
    page.fill("[data-testid=el-html]", "<p class='custom'>raw</p>")
    page.select_option("[data-testid=el-fragment]", "grow")
    frame = page.frame_locator("[data-testid=preview]")
    expect(frame.locator("section.present p.custom")).to_have_text("raw")
    page.click("[data-testid=add-slide] summary")
    page.click("[data-testid=add-slide-markdown]")
    page.click("[data-testid=tab-slide]")
    page.fill("[data-testid=slide-markdown]", "## From Markdown\n\n- **bold** point")
    frame = page.frame_locator("[data-testid=preview]")
    expect(frame.locator("section.present h2")).to_have_text("From Markdown")
    expect(frame.locator("section.present strong")).to_have_text("bold")
    text = save_html(page, tmp_path).read_text(encoding="utf-8")
    assert '<div class="fragment grow"><p class=\'custom\'>raw</p></div>' in text or "fragment grow" in text
    assert "<section data-markdown" in text and "## From Markdown" in text
    assert "<table" in text and "<blockquote" in text
    assert page.errors == []


def test_import_markdown_notebook_and_reveal_html(app_page, tmp_path):
    page = ready(app_page)
    md = tmp_path / "talk.md"
    md.write_text("# Markdown talk\n\nintro\n\n---\n\n## Part 2\n\nNote:\nsecret notes\n\n--\n\n## Detail\n", encoding="utf-8")
    open_file(page, md)
    expect(page.locator("[data-testid=slide-item]")).to_have_count(3)
    expect(page.locator("[data-testid=slide-item].sub")).to_have_count(1)
    assert page.evaluate("RV.state.deck.slides[1].notes") == "secret notes"
    assert page.evaluate("RV.state.name") == "talk.html"

    nb = tmp_path / "analysis.ipynb"
    nb.write_text(json.dumps({"nbformat": 4, "nbformat_minor": 5, "metadata": {"language_info": {"name": "python"}}, "cells": [
        {"cell_type": "markdown", "metadata": {"slideshow": {"slide_type": "slide"}}, "source": ["# Results"]},
        {"cell_type": "code", "metadata": {"slideshow": {"slide_type": "fragment"}}, "execution_count": 1, "source": ["1 + 1"],
         "outputs": [{"output_type": "execute_result", "execution_count": 1, "metadata": {}, "data": {"text/plain": ["2"]}}]},
        {"cell_type": "markdown", "metadata": {"slideshow": {"slide_type": "skip"}}, "source": ["hidden"]},
        {"cell_type": "markdown", "metadata": {"slideshow": {"slide_type": "subslide"}}, "source": ["## Details"]},
        {"cell_type": "markdown", "metadata": {"slideshow": {"slide_type": "notes"}}, "source": ["say this"]},
    ]}), encoding="utf-8")
    open_file(page, nb)
    expect(page.locator("[data-testid=slide-item]")).to_have_count(2)
    deck = page.evaluate("RV.state.deck")
    assert "```python\n1 + 1\n```" in deck["slides"][0]["markdown"] and 'class="fragment"' in deck["slides"][0]["markdown"]
    assert "hidden" not in json.dumps(deck) and deck["slides"][0]["children"][0]["notes"] == "say this"
    frame = page.frame_locator("[data-testid=preview]")
    expect(frame.locator("section.present h1")).to_have_text("Results")

    foreign = tmp_path / "foreign.html"
    foreign.write_text("""<html><head><title>Foreign</title><link rel="stylesheet" href="dist/theme/serif.css"></head><body>
      <div class="reveal"><div class="slides">
        <section data-background-color="#ff0000"><h2>One</h2><aside class="notes">n1</aside></section>
        <section><section><h2>Two</h2></section><section data-transition="fade"><h2>Two b</h2></section></section>
      </div></div><script>Reveal.initialize({ controls: false, transition: 'convex', width: 1280, height: 720 });</script></body></html>""", encoding="utf-8")
    open_file(page, foreign)
    expect(page.locator("[data-testid=slide-item]")).to_have_count(3)
    deck = page.evaluate("RV.state.deck")
    assert deck["theme"] == "serif" and deck["config"]["transition"] == "convex" and deck["config"]["controls"] is False
    assert deck["config"]["width"] == 1280 and deck["slides"][0]["background"]["color"] == "#ff0000" and deck["slides"][0]["notes"] == "n1"
    assert deck["slides"][1]["children"][0]["transition"] == "fade"
    assert page.errors == []


def test_present_mode(app_page):
    page = ready(app_page)
    page.click("[data-testid=present]")
    expect(page.locator("[data-testid=presenter]")).to_be_visible()
    presenter = page.frame_locator("[data-testid=presenter-frame]")
    expect(presenter.locator(".reveal.ready")).to_have_count(1, timeout=10000)
    page.locator("[data-testid=presenter-frame]").focus()
    page.keyboard.press("ArrowRight")
    expect(presenter.locator("section.present h2")).to_have_text("Heading")
    page.click("[data-testid=present-close]")
    expect(page.locator("[data-testid=presenter]")).to_be_hidden()


def test_notebook_slide_types_open_in_web_slides(app_page, tmp_path):
    """The notebook marks cells as slides and opens them in the reveal.js editor ("Open with" bridge)."""
    page = app_page("notebook", init_script=NO_PICKERS)
    page.click("[data-testid=slideshow]")
    page.locator("[data-testid=slide-type]").first.select_option("slide")
    with page.context.expect_page() as new_page:
        page.click("[data-testid=as-slides]")
    slides = new_page.value
    slides.wait_for_selector("body[data-ready]", timeout=30000)
    expect(slides.locator("[data-testid=slide-item]").first).to_contain_text("Python notebook")
    assert slides.evaluate("RV.state.deck.slides.length") >= 1


def test_mobile_layout(app_page):
    page = ready(app_page, mobile=True)
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
    expect(page.locator("[data-testid=inspector]")).to_be_hidden()
    page.click("[data-testid=phone-next]")
    expect(page.locator("[data-testid=slide-item].active")).to_have_count(1)
    assert page.evaluate("[RV.state.h, RV.state.v]") == [1, 0]
    page.click("[data-testid=toggle-panels]")
    expect(page.locator("[data-testid=inspector]")).to_be_visible()
    page.click("[data-testid=phone-slides]")
    expect(page.locator("[data-testid=slide-list]")).to_be_visible()
    assert page.errors == []


def test_german_interface(app_page):
    page = app_page("reveal", init_script=NO_PICKERS + "localStorage.setItem('suite.lang', 'de');")
    expect(page.locator("[data-testid=present]")).to_contain_text("Präsentieren")
    expect(page.locator("[data-testid=tab-deck]")).to_have_text("Präsentation")


MATH_RENDERED = "() => { const d = document.querySelector('[data-testid=preview]').contentDocument; return !!(d && d.body && d.body.dataset.math === 'rendered'); }"


def test_math(app_page, tmp_path, browser):
    page = ready(app_page)
    # no formulas: no KaTeX in the saved file
    assert "katex" not in save_html(page, tmp_path, "plain.html").read_text(encoding="utf-8").lower()
    # a Markdown slide with a display formula and an inline formula; a single $ stays text
    page.click("[data-testid=add-slide] summary")
    page.click("[data-testid=add-slide-markdown]")
    page.click("[data-testid=tab-slide]")
    v = preview_version(page)
    page.fill("[data-testid=slide-markdown]", "## Energy\n\n$$E = mc^2$$\n\nInline \\(a^2 + b^2\\) costs $5")
    frame = wait_preview(page, v)
    page.wait_for_function(MATH_RENDERED, timeout=15000)
    expect(frame.locator("section.present .katex-display")).to_have_count(1)
    expect(frame.locator("section.present .katex")).to_have_count(2)
    expect(frame.locator("section.present")).to_contain_text("costs $5")
    # editing a text with a formula edits its TeX source
    page.click("[data-testid=slide-item] >> nth=0")
    v = preview_version(page)
    page.evaluate("() => { const s = RV.state.deck.slides[0]; s.blocks[0].html = 'Area \\\\(\\\\pi r^2\\\\)'; RV.updatePreview(); }")
    frame = wait_preview(page, v)
    page.wait_for_function(MATH_RENDERED, timeout=15000)
    heading = frame.locator("section.present [data-block]").first
    expect(heading.locator(".katex")).to_have_count(1)
    heading.dblclick()
    expect(page.locator("[data-testid=format-bar]")).to_be_visible()
    expect(heading).to_contain_text("\\(\\pi r^2\\)")
    page.click("[data-testid=fmt-done]")
    assert page.evaluate("RV.state.deck.slides[0].blocks[0].html") == "Area \\(\\pi r^2\\)"
    # the saved presentation typesets offline
    path = save_html(page, tmp_path, "math.html")
    text = path.read_text(encoding="utf-8")
    assert "renderMathInElement" in text and "data:font/woff2" in text
    assert "Inline \\\\(a^2 + b^2\\\\)" in text  # escaped for the Markdown plugin
    # a presentation from elsewhere (no editor model) is read back without the escapes
    foreign = text.replace('id="suite-reveal-deck"', 'id="other"')
    md = page.evaluate("(t) => RV.fromHtml(t).slides.map((s) => s.markdown).join('|')", foreign)
    assert "Inline \\(a^2 + b^2\\) costs $5" in md
    context = browser.new_context(offline=True)
    show = context.new_page()
    show.goto(path.as_uri())
    show.wait_for_function("document.body.dataset.math === 'rendered'", timeout=15000)
    assert show.locator(".katex").count() >= 3
    context.close()
    assert page.errors == []


def test_powerpoint_to_web_slides_and_back(app_page, tmp_path):
    import io
    from pptx import Presentation
    from helpers import office_samples
    source = tmp_path / "bericht.pptx"
    office_samples.make_pptx(source)
    page = ready(app_page)
    open_file(page, source)
    expect(page.locator("[data-testid=toast]").last).to_contain_text("converted into web slides")
    slides = page.evaluate("RV.state.deck.slides.map((s) => s.blocks.map((b) => [b.type, b.html || (b.items || []).join('|') || (b.rows ? JSON.stringify(b.rows) : '') || (b.src ? 'image' : '')]))")
    assert slides[0][0] == ["heading", "Quartalsbericht 2026"] and slides[0][1][1].startswith("Vertraulich")
    assert slides[1][0] == ["heading", "Agenda"] and slides[1][1][0] == "list" and "Inland" in slides[1][1][1] and "&nbsp;" in slides[1][1][1]
    third = {}
    for kind, v in slides[2]:
        third[kind] = third.get(kind, "") + v
    assert "<strong>fett</strong>" in third["text"] and "color:#FF0000" in third["text"] and third["image"] == "image" and "Äpfel" in third["table"]
    assert page.evaluate("RV.state.deck.slides[2].background.color") == "#FFF2CC"
    assert page.evaluate("RV.state.name") == "bericht.html"
    # and saved as PowerPoint again
    page.click("[data-testid=save-as]")
    page.select_option("[data-testid=save-format]", "pptx")
    expect(page.locator("[data-testid=pptx-hint]")).to_be_visible()
    with page.expect_download() as d:
        page.click("[data-testid=save-ok]")
    assert d.value.suggested_filename == "bericht.pptx"
    prs = Presentation(io.BytesIO(open(d.value.path(), "rb").read()))
    assert len(prs.slides) == 3
    texts = [[sh.text_frame.text for sh in s.shapes if sh.has_text_frame] for s in prs.slides]
    assert texts[0][0] == "Quartalsbericht 2026" and "Inland" in "\n".join(texts[1])
    kinds = [sh.shape_type for sh in prs.slides[2].shapes]
    assert any(sh.has_table for sh in prs.slides[2].shapes if hasattr(sh, "has_table")) and 13 in kinds  # a table and a picture (MSO_SHAPE_TYPE.PICTURE)
    assert page.errors == []


def test_diagram_block_drag_and_drop_and_print_with_notes(app_page, tmp_path):
    page = ready(app_page)
    before = preview_version(page)
    page.click("[data-testid=add-block] summary")
    page.click("[data-testid=add-block-diagram]")
    frame = wait_preview(page, before)
    expect(frame.locator(".suite-diagram svg")).to_have_count(1)
    expect(frame.locator(".suite-diagram svg text")).to_contain_text(["Idea", "Good?", "Slides"])
    # an own flowchart; an error is shown and the last good drawing kept out of the slide
    page.fill("[data-testid=el-mermaid]", "graph TD\n  X[Start] --> Y[Ende]\n  X --> Z[Auch]")
    expect(frame.locator(".suite-diagram svg text")).to_have_text(["Start", "Ende", "Auch"])
    page.fill("[data-testid=el-mermaid]", "sequenceDiagram\n  A->>B: hi")
    expect(page.locator("[data-testid=el-mermaid-error]")).to_contain_text("flowcharts")
    page.fill("[data-testid=el-mermaid]", "graph LR\n  P[Plan] --> Q[Tun]")
    expect(frame.locator(".suite-diagram svg text")).to_have_text(["Plan", "Tun"])
    # saved as part of the presentation
    path = save_html(page, tmp_path)
    assert path.read_text(encoding="utf-8").count("<svg") >= 1 and ">Plan</tspan>" in path.read_text(encoding="utf-8")

    # drag and drop in the slide list: slide 2 before slide 1, then slide 1 into slide 2 as a sub-slide
    page.click("[data-testid=add-slide] summary")
    page.click("[data-testid=add-slide-quote]")
    items = page.locator("[data-testid=slide-item]")
    expect(items).to_have_count(3)
    titles = lambda: page.evaluate("RV.state.deck.slides.map((s) => [s.blocks[0] ? s.blocks[0].type : '', s.children.length])")
    assert [t for t, _ in titles()] == ["heading", "quote", "heading"]  # inserted after the current slide
    touch = page.evaluate("matchMedia('(pointer: coarse)').matches")  # phones use the ↑ ↓ ⇥ ⇤ buttons instead
    if not touch:
        drag_slides(page, items, titles)
    else:
        print("drag and drop skipped on touch screens")

    # print with speaker notes: reveal.js print view with the notes on separate pages
    print_with_notes(page)
    assert page.errors == []


def drag_slides(page, items, titles):
    items.nth(1).drag_to(items.nth(0), target_position={"x": 10, "y": 3})
    expect(items.first.locator(".rv-num")).to_have_text("1")
    assert [t for t, _ in titles()] == ["quote", "heading", "heading"]
    box = items.nth(1).bounding_box()
    items.nth(0).drag_to(items.nth(1), target_position={"x": box["width"] - 8, "y": box["height"] / 2})
    expect(page.locator("[data-testid=slide-item].sub")).to_have_count(1)
    assert titles() == [["heading", 1], ["heading", 0]]
    page.click("[data-testid=undo]")
    assert len(titles()) == 3


def print_with_notes(page):
    page.click("[data-testid=print] summary")
    with page.expect_popup() as popup:
        page.click("[data-testid=print-notes]")
    win = popup.value
    win.wait_for_load_state()
    content = win.content()
    assert '"showNotes":"separate-page"' in content.replace(" ", "") and '"view":"print"' in content.replace(" ", "")
    win.close()
