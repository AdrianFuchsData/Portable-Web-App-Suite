"""Spike S-11 (docs/planning/spike-results.md): a PNG that carries a draw.io diagram (tEXt chunk "mxfile") keeps it through
the PPTX round trip of the presentation editor and the DOCX round trip of the document editor."""
import base64
import io
import struct
import urllib.parse
import zipfile
import zlib

from PIL import Image
from playwright.sync_api import expect

NO_PICKERS = "window.showOpenFilePicker = undefined; window.showSaveFilePicker = undefined;"
XML = ('<mxfile><diagram name="P"><mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/><mxCell id="a" value="Spike" style="rounded=1;" '
       'vertex="1" parent="1"><mxGeometry x="10" y="10" width="120" height="60" as="geometry"/></mxCell></root></mxGraphModel></diagram></mxfile>')


def png_with_diagram():
    buf = io.BytesIO()
    Image.new("RGB", (140, 80), "white").save(buf, "PNG")
    png = buf.getvalue()
    data = b"mxfile\0" + urllib.parse.quote(XML, safe="").encode()
    chunk = struct.pack(">I", len(data)) + b"tEXt" + data + struct.pack(">I", zlib.crc32(b"tEXt" + data))
    return png[:33] + chunk + png[33:]  # after IHDR


def has_diagram(png):
    p = 8
    while p + 8 <= len(png):
        n, typ = struct.unpack(">I4s", png[p:p + 8])
        if typ == b"tEXt" and png[p + 8:p + 8 + n].startswith(b"mxfile\0"):
            return urllib.parse.unquote(png[p + 15:p + 8 + n].decode()) == XML
        p += 12 + n
    return False


def media_pngs(path, folder):
    with zipfile.ZipFile(path) as z:
        return [z.read(n) for n in z.namelist() if n.startswith(folder) and n.endswith(".png")]


def from_data_url(src):
    return base64.b64decode(src.split(",", 1)[1] + "==") if src.startswith("data:image/png") else b""


def test_presentation_pptx_keeps_the_diagram(app_page, tmp_path):
    page = app_page("presentation", init_script=NO_PICKERS)
    src = "data:image/png;base64," + base64.b64encode(png_with_diagram()).decode()
    page.evaluate("""(src) => { PR.state.deck.slides[0].elements.push({ id: 'dia', type: 'image', name: 'diagram.png', x: 100, y: 100, w: 280, h: 160, rot: 0,
      geometry: 'rect', fill: null, line: null, text: null, image: { src } }); PR.render(); }""", src)
    with page.expect_download() as d:
        page.click("[data-testid=save]")
    pptx = tmp_path / "spike.pptx"
    d.value.save_as(pptx)
    assert any(has_diagram(x) for x in media_pngs(pptx, "ppt/media/"))
    page.evaluate("(b) => PR.openBase64(b, 'spike.pptx')", base64.b64encode(pptx.read_bytes()).decode())
    page.wait_for_function("() => PR.state.name === 'spike.pptx'")
    srcs = page.evaluate("PR.state.deck.slides.flatMap((s) => s.elements).filter((e) => e.type === 'image').map((e) => e.image.src)")
    assert any(has_diagram(from_data_url(s)) for s in srcs)
    assert page.errors == []


def test_document_docx_keeps_the_diagram(app_page, tmp_path):
    import docx
    from docx.shared import Cm
    source = docx.Document()
    source.add_paragraph("Before")
    source.add_picture(io.BytesIO(png_with_diagram()), width=Cm(6))
    source.add_paragraph("After")
    path = tmp_path / "in.docx"
    source.save(path)
    page = app_page("document", init_script=NO_PICKERS)
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(path))
    expect(page.locator("[data-testid=file-name]")).to_contain_text("in.docx", timeout=20000)
    page.wait_for_function("() => Object.keys(DOC.doc.save().drawings || {}).length === 1", timeout=20000)
    sources = page.evaluate("Object.values(DOC.doc.save().drawings).map((d) => d.source || '')")
    assert any(has_diagram(from_data_url(s)) for s in sources)
    page.evaluate("DOC.state.dirty = true")
    with page.expect_download() as d:
        page.click("[data-testid=save]")
    out = tmp_path / "out.docx"
    d.value.save_as(out)
    assert any(has_diagram(x) for x in media_pngs(out, "word/media/"))
    assert page.errors == []
