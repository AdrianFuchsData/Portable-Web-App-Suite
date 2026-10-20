"""Spreadsheet (apps/spreadsheet) on Univer Sheets. The file conversion itself is tested in tests/prototypes/test_s6_office_io.py."""
import csv
import io
import zipfile

import openpyxl
from playwright.sync_api import expect

from helpers import office_samples

NO_PICKERS = "window.showOpenFilePicker = undefined; window.showSaveFilePicker = undefined;"
VALUE = "([a1, sheet]) => { const wb = SS.workbook; const ws = sheet ? wb.getSheetByName(sheet) : wb.getActiveSheet(); return ws.getRange(a1).getValue(); }"


def open_file(page, path):
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(path))
    expect(page.locator("[data-testid=file-name]")).to_have_text(path.name, timeout=30000)
    page.wait_for_timeout(1500)  # formulas are calculated asynchronously


def value(page, a1, sheet=None):
    return page.evaluate(VALUE, [a1, sheet])


def ready(app_page, **kw):
    page = app_page("spreadsheet", init_script=NO_PICKERS, **kw)
    page.wait_for_selector("[data-testid=sheet] canvas", timeout=60000)
    page.wait_for_timeout(1600)  # the first workbook is set up asynchronously
    return page


def test_open_edit_and_save_xlsx(app_page, tmp_path):
    source = office_samples.make_xlsx(tmp_path / "sample.xlsx")
    page = ready(app_page)
    assert page.title() == "Spreadsheet"
    expect(page.locator("[data-testid=file-name]")).not_to_contain_text("•")
    open_file(page, source)
    page.wait_for_timeout(500)
    assert value(page, "D5") == 63.5
    page.evaluate("SS.workbook.getActiveSheet().getRange('B2').activate()")
    page.keyboard.type("10")
    page.keyboard.press("Enter")
    page.wait_for_timeout(1000)
    assert value(page, "D5") == 74  # D2 = 10 * 1.5 instead of 3 * 1.5: the sum is recalculated
    expect(page.locator("[data-testid=file-name]")).to_contain_text("•")
    with page.expect_download() as d:
        page.click("[data-testid=save]")
    assert d.value.suggested_filename == "sample.xlsx"
    expect(page.locator("[data-testid=banner]")).to_contain_text("sample.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(open(d.value.path(), "rb").read()))
    ws = wb["Data"]
    assert ws["B2"].value == 10 and ws["D2"].value == "=B2*C2" and ws["A1"].font.bold
    expect(page.locator("[data-testid=file-name]")).not_to_contain_text("•")
    assert page.errors == [] and page.requests == []


def test_typing_formula_in_new_workbook(app_page):
    page = ready(app_page)
    page.evaluate("SS.workbook.getActiveSheet().getRange('A1').activate()")
    page.keyboard.type("21")
    page.keyboard.press("Enter")
    page.evaluate("SS.workbook.getActiveSheet().getRange('A2').activate()")
    page.keyboard.type("=A1*2")
    page.keyboard.press("Enter")
    page.wait_for_timeout(1000)
    assert value(page, "A2") == 42
    # undo/redo from Univer's toolbar (also usable on phones without a keyboard)
    page.click("[data-u-command='univer.command.undo']")
    page.wait_for_timeout(500)
    assert value(page, "A2") is None
    page.click("[data-u-command='univer.command.redo']")
    page.wait_for_timeout(500)
    assert value(page, "A2") == 42


def test_save_as_other_formats(app_page, tmp_path):
    source = office_samples.make_xlsx(tmp_path / "sample.xlsx")
    page = ready(app_page)
    open_file(page, source)
    for fmt in ("ods", "csv", "xls"):
        page.click("[data-testid=save-as]")
        page.select_option("[data-testid=save-as-format]", fmt)
        with page.expect_download() as d:
            page.click("[data-testid=save-as-ok]")
        assert d.value.suggested_filename == f"sample.{fmt}"
        data = open(d.value.path(), "rb").read()
        if fmt == "ods":
            assert zipfile.ZipFile(io.BytesIO(data)).read("mimetype") == b"application/vnd.oasis.opendocument.spreadsheet"
        elif fmt == "csv":
            rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
            assert rows[0][:4] == ["Item", "Qty", "Price", "Total"] and rows[1][0] == "Äpfel"
        else:
            assert data[:8] == bytes.fromhex("D0CF11E0A1B11AE1")  # BIFF8 in an OLE container
        expect(page.locator("[data-testid=file-name]")).to_have_text(f"sample.{fmt}")
    expect(page.locator("[data-testid=toast]").last).to_contain_text("only values and formulas")


def test_xlsm_is_saved_as_xlsx(app_page, tmp_path):
    xlsm = tmp_path / "macro.xlsm"
    office_samples.make_xlsx(xlsm)
    page = ready(app_page)
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(xlsm))
    expect(page.locator("[data-testid=file-name]")).to_have_text("macro.xlsx", timeout=30000)
    expect(page.locator("[data-testid=toast]").last).to_contain_text("Macros cannot be kept")


def test_new_discards_after_confirmation(app_page, tmp_path):
    page = ready(app_page)
    open_file(page, office_samples.make_xlsx(tmp_path / "sample.xlsx"))
    page.evaluate("SS.workbook.getActiveSheet().getRange('A20').activate()")
    page.locator("[data-testid=sheet] canvas").first.click(position={"x": 5, "y": 5})  # user input in the sheet
    page.evaluate("SS.workbook.getActiveSheet().getRange('A20').activate()")
    page.keyboard.type("changed")
    page.keyboard.press("Enter")
    expect(page.locator("[data-testid=file-name]")).to_contain_text("•")
    page.click("[data-testid=new]")
    page.click("[data-testid=confirm-cancel]")
    assert value(page, "A20") == "changed"
    page.click("[data-testid=new]")
    page.click("[data-testid=confirm-ok]")
    expect(page.locator("[data-testid=file-name]")).to_have_text("Spreadsheet.xlsx")


def test_language_and_theme(app_page):
    page = ready(app_page)
    page.click("[data-testid=language]")
    expect(page.locator("[data-testid=app-title]")).to_have_text("Tabellenkalkulation")
    expect(page.get_by_text("Formeln", exact=True)).to_be_visible()  # Univer's own UI follows the language
    page.click("[data-testid=theme]")
    page.click("[data-testid=theme]")
    assert page.errors == []


def test_mobile_layout(app_page):
    page = ready(app_page, mobile=True)
    expect(page.locator("[data-testid=sheet] canvas").first).to_be_visible()
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
    assert page.errors == []


def test_csv_import_dialog(app_page, tmp_path):
    source = tmp_path / "export.csv"
    text = 'Name;PLZ;Betrag;Datum;Notiz\r\nMüller;0815;1.234,50;16.10.2026;"a;b\r\nzweite Zeile"\r\nÖzil;4020;-7,5;2026-01-02;x\r\n'
    source.write_bytes(text.encode("windows-1252"))
    page = ready(app_page)
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(source))
    dialog = page.locator("[data-testid=csv-dialog]")
    expect(dialog).to_be_visible()
    # detected: Windows-1252, semicolon, decimal comma
    expect(dialog.locator("[data-testid=csv-encoding] option:checked")).to_have_text("Automatic (WINDOWS-1252)")
    expect(dialog.locator("[data-testid=csv-separator]")).to_have_value(";")
    expect(dialog.locator("[data-testid=csv-decimal]")).to_have_value(",")
    preview = dialog.locator("[data-testid=csv-preview] tr")
    expect(preview).to_have_count(3)
    expect(preview.nth(1).locator("td")).to_have_text(["Müller", "0815", "1234.5", "2026-10-16", "a;b\r\nzweite Zeile"])
    # a wrong choice changes the preview
    dialog.locator("[data-testid=csv-separator]").select_option(",")
    expect(preview.first.locator("td").first).to_have_text("Name;PLZ;Betrag;Datum;Notiz")
    dialog.locator("[data-testid=csv-separator]").select_option(";")
    page.click("[data-testid=csv-import]")
    expect(page.locator("[data-testid=file-name]")).to_have_text("export.csv", timeout=30000)
    page.wait_for_timeout(1000)
    assert value(page, "A2") == "Müller" and value(page, "A3") == "Özil"
    assert value(page, "B2") == "0815"  # leading zeros stay text
    assert value(page, "C2") == 1234.5 and value(page, "C3") == -7.5
    assert value(page, "E2") == "a;b\nzweite Zeile"
    assert value(page, "D2") == "2026-10-16" and value(page, "D3") == "2026-01-02"  # date cells (serial number with a date format)
    assert page.evaluate("SS.workbook.getActiveSheet().getRange('D2').getCellData().t") == 2  # a number
    # saved again as CSV with the same separator
    with page.expect_download() as d:
        page.click("[data-testid=save]")
    rows = list(csv.reader(io.StringIO(open(d.value.path(), "rb").read().decode("utf-8-sig")), delimiter=";"))
    assert rows[0] == ["Name", "PLZ", "Betrag", "Datum", "Notiz"] and rows[1][:3] == ["Müller", "0815", "1234.5"]
    # everything as text
    with page.expect_file_chooser() as chooser:
        page.click("[data-testid=open]")
    chooser.value.set_files(str(source))
    page.select_option("[data-testid=csv-types]", "text")
    expect(page.locator("[data-testid=csv-decimal]")).to_have_count(0)
    page.click("[data-testid=csv-import]")
    page.wait_for_timeout(1500)
    assert value(page, "C2") == "1.234,50" and value(page, "D2") == "16.10.2026"
    assert page.errors == []


def make_rules_xlsx(path):
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daten"
    for i in range(1, 11):
        ws.cell(i, 1, i)
        ws.cell(i, 2, i * 10)
        ws.cell(i, 3, "x" if i % 2 else "abc")
    ws.conditional_formatting.add("A1:A10", CellIsRule(operator="greaterThan", formula=["5"], fill=PatternFill(bgColor="FFFF0000", fill_type="solid"), font=Font(color="FFFFFFFF", bold=True)))
    ws.conditional_formatting.add("B1:B10", ColorScaleRule(start_type="min", start_color="FFF8696B", end_type="max", end_color="FF63BE7B"))
    ws.conditional_formatting.add("C1:C10", FormulaRule(formula=['ISNUMBER(SEARCH("ab",C1))'], fill=PatternFill(bgColor="FFFFFF00", fill_type="solid")))
    ws.conditional_formatting.add("D1:D10", DataBarRule(start_type="min", end_type="max", color="FF638EC6"))
    dv = DataValidation(type="list", formula1='"ja,nein,vielleicht"', allow_blank=True, errorStyle="warning", error="Bitte wählen", errorTitle="Ungültig", prompt="Auswahl", promptTitle="Hinweis")
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add("E1:E5")
    dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="10")
    ws.add_data_validation(dv2)
    dv2.add("F1:F3")
    wb.save(path)
    return path


def test_conditional_formatting_and_validation(app_page, tmp_path):
    page = ready(app_page)
    open_file(page, make_rules_xlsx(tmp_path / "regeln.xlsx"))
    cf = page.evaluate("SS.workbook.getActiveSheet().getConditionalFormattingRules().map((r) => ({ type: r.rule.type, sub: r.rule.subType, op: r.rule.operator, value: r.rule.value, style: r.rule.style, n: r.ranges.length, row: r.ranges[0].endRow }))")
    assert [(r["type"], r["sub"]) for r in cf] == [("highlightCell", "number"), ("colorScale", None), ("highlightCell", "formula"), ("dataBar", None)]
    assert cf[0]["op"] == "greaterThan" and cf[0]["value"] == 5 and cf[0]["style"]["bg"]["rgb"].lower() == "#ff0000" and cf[0]["style"]["bl"] == 1 and cf[0]["row"] == 9
    assert cf[2]["value"] == '=ISNUMBER(SEARCH("ab",C1))'
    dv = page.evaluate("SS.workbook.getActiveSheet().getDataValidations().map((d) => { const r = d.rule; return { type: r.type, f1: r.formula1, f2: r.formula2, op: r.operator, ranges: r.ranges.map((x) => [x.startRow, x.endRow, x.startColumn]), error: r.error, style: r.errorStyle }; })")
    by_type = {d["type"]: d for d in dv}
    assert by_type["list"]["f1"] == "ja,nein,vielleicht" and by_type["list"]["ranges"] == [[0, 4, 4]] and by_type["list"]["error"] == "Bitte wählen" and by_type["list"]["style"] == 2
    assert by_type["whole"]["op"] == "between" and (by_type["whole"]["f1"], by_type["whole"]["f2"]) == ("1", "10")
    # a new rule of each kind, then save
    page.evaluate("""() => { const api = SS.univerAPI; const ws = SS.workbook.getActiveSheet();
      ws.addConditionalFormattingRule(ws.newConditionalFormattingRule().whenTextContains('ab').setBackground('#00ff00').setRanges([ws.getRange('G1:G5').getRange()]).build());
      ws.getRange('H1:H4').setDataValidation(api.newDataValidation().requireNumberBetween(0, 100).build()); }""")
    page.keyboard.press("Control+s")
    page.wait_for_timeout(500)
    with page.expect_download() as d:
        page.click("[data-testid=save]")
    wb = openpyxl.load_workbook(io.BytesIO(open(d.value.path(), "rb").read()))
    ws = wb["Daten"]
    rules = {str(rng.sqref): [r for r in rng.rules] for rng in ws.conditional_formatting}
    assert rules["A1:A10"][0].type == "cellIs" and rules["A1:A10"][0].operator == "greaterThan" and rules["A1:A10"][0].formula == ["5"]
    assert rules["A1:A10"][0].dxf.fill.bgColor.rgb.endswith("FF0000") and rules["A1:A10"][0].dxf.font.b
    assert rules["B1:B10"][0].type == "colorScale" and rules["D1:D10"][0].type == "dataBar"
    assert rules["C1:C10"][0].type == "expression" and rules["C1:C10"][0].formula == ['ISNUMBER(SEARCH("ab",C1))']
    assert rules["G1:G5"][0].type == "containsText" and rules["G1:G5"][0].formula == ['NOT(ISERROR(SEARCH("ab",G1)))']
    validations = {(v.type, str(v.sqref)): v for v in ws.data_validations.dataValidation}
    lst = next(v for (t, _), v in validations.items() if t == "list")
    assert lst.formula1 == '"ja,nein,vielleicht"' and lst.errorStyle == "warning" and lst.error == "Bitte wählen" and str(lst.sqref) in ("E1:E5", "E1 E2 E3 E4 E5")
    whole = [v for (t, _), v in validations.items() if t == "whole"]
    assert whole and whole[0].formula1 == "1" and whole[0].formula2 == "10"
    dec = [v for (t, _), v in validations.items() if t == "decimal"]
    assert dec and dec[0].formula1 == "0" and dec[0].formula2 == "100" and "H1" in str(dec[0].sqref)
    # and opened again: the text rule is complete
    path = tmp_path / "wieder.xlsx"
    d.value.save_as(path)
    open_file(page, path)
    texts = page.evaluate("SS.workbook.getActiveSheet().getConditionalFormattingRules().filter((r) => r.rule.subType === 'text').map((r) => [r.rule.operator, r.rule.value])")
    assert texts == [["containsText", "ab"]]
    assert page.errors == []
