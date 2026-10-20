"""Sample office files for the import/export tests (written with openpyxl, converted with LibreOffice)."""
import datetime
import re
import shutil
import subprocess
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SOFFICE = shutil.which("soffice")

# What the sample contains (used by the tests)
DATA_ROWS = [("Äpfel", 3, 1.5, datetime.date(2026, 1, 15), True),
             ("Birnen", 4, 2.25, datetime.date(2026, 2, 28), False),
             ("Kirschen 🍒", 5, 10, datetime.date(2026, 12, 31), True)]


def make_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Qty", "Price", "Total", "Date", "Ok"])
    for i, (item, qty, price, date, ok) in enumerate(DATA_ROWS, start=2):
        ws.append([item, qty, price, f"=B{i}*C{i}", date, ok])
    ws["A5"] = "Sum"
    ws["D5"] = "=SUM(D2:D4)"
    ws["A7"] = "Merged title"
    ws.merge_cells("A7:D7")

    header = ws["A1"]
    header.font = Font(bold=True, color="FF0000FF")
    header.fill = PatternFill("solid", fgColor="FFFFFF00")
    header.border = Border(bottom=Side(style="thin", color="FF000000"))
    header.alignment = Alignment(horizontal="center", vertical="top")
    ws["B1"].alignment = Alignment(vertical="center")
    ws["A5"].font = Font(italic=True, color="FFFF0000")
    ws["D5"].border = Border(top=Side(style="double", color="FF00FF00"))
    ws["D5"].font = Font(bold=True, underline="single", strike=True)
    for row in range(2, 5):
        ws[f"C{row}"].number_format = "#,##0.00"
        ws[f"E{row}"].number_format = "yyyy-mm-dd"
    title = ws["A7"]
    title.font = Font(name="Arial", size=16)
    title.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 8
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    second = wb.create_sheet("Second")
    second["A1"] = "=Data!D5*2"
    second["A2"] = "Text with umlauts: äöü ß"
    second.sheet_properties.tabColor = "FF00B050"
    wb.save(path)
    _make_shared_formulas(path)
    return path


def _make_shared_formulas(path):
    """Rewrite D2:D4 of the first sheet as one shared formula, as Excel writes it."""
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    xml = files["xl/worksheets/sheet1.xml"].decode()
    xml = xml.replace("<f>B2*C2</f>", '<f t="shared" ref="D2:D4" si="0">B2*C2</f>', 1)
    xml = re.sub(r"<f>B[34]\*C[34]</f>", '<f t="shared" si="0"/>', xml)
    assert xml.count('t="shared"') == 3, xml
    files["xl/worksheets/sheet1.xml"] = xml.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)


def convert(path, target_format, outdir):
    """Convert with LibreOffice (headless). Returns the path of the converted file."""
    profile = outdir / "lo-profile"
    subprocess.run([SOFFICE, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", target_format,
                    "--outdir", str(outdir), str(path)], check=True, capture_output=True, timeout=180)
    out = outdir / (path.stem + "." + target_format.split(":")[0])
    assert out.exists(), f"LibreOffice did not create {out}"
    return out


def make_docx(path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading("Projektbericht", level=1)
    p = doc.add_paragraph("Normal text with ")
    p.add_run("bold").bold = True
    p.add_run(", ")
    p.add_run("italic").italic = True
    p.add_run(", ")
    p.add_run("underlined").underline = True
    p.add_run(" and ")
    red = p.add_run("red")
    red.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    red.font.size = Pt(14)
    p.add_run(".")
    doc.add_heading("Liste", level=2)
    for item in ("Erster Punkt", "Zweiter Punkt", "Dritter Punkt"):
        doc.add_paragraph(item, style="List Bullet")
    for item in ("Schritt eins", "Schritt zwei"):
        doc.add_paragraph(item, style="List Number")
    table = doc.add_table(rows=2, cols=3)
    for c, text in enumerate(("Name", "Menge", "Preis")):
        table.cell(0, c).text = text
    for c, text in enumerate(("Äpfel", "3", "1,50")):
        table.cell(1, c).text = text
    centered = doc.add_paragraph("Zentriert äöü ß 🍒")
    centered.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Second page")
    doc.save(path)
    return path


def make_pptx(path):
    """Three slides: title layout (inherited placeholders), bullets, and free content."""
    import io as _io

    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()  # 4:3, 10 x 7.5 in
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Quartalsbericht 2026"
    s1.placeholders[1].text = "Vertraulich – äöü"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Agenda"
    body = s2.placeholders[1].text_frame
    body.text = "Umsatz"
    for text, level in (("Inland", 1), ("Export", 1), ("Kosten", 0)):
        p = body.add_paragraph()
        p.text = text
        p.level = level

    s3 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    s3.background.fill.solid()
    s3.background.fill.fore_color.rgb = RGBColor(0xFF, 0xF2, 0xCC)
    box = s3.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(1))
    p = box.text_frame.paragraphs[0]
    p.add_run().text = "Normal "
    bold = p.add_run()
    bold.text = "fett"
    bold.font.bold = True
    p.add_run().text = " "
    red = p.add_run()
    red.text = "kursiv rot"
    red.font.italic = True
    red.font.size = Pt(28)
    red.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    rect = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(2), Inches(3), Inches(1.5))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0x44, 0x72, 0xC4)
    rect.text_frame.text = "Box"
    oval = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(2), Inches(2), Inches(1.5))
    oval.fill.background()
    oval.line.color.rgb = RGBColor(0x00, 0x80, 0x00)
    oval.line.width = Pt(3)
    image = _io.BytesIO()
    Image.new("RGB", (200, 100), (200, 30, 30)).save(image, "PNG")
    image.seek(0)
    s3.shapes.add_picture(image, Inches(6.5), Inches(2), Inches(2), Inches(1))
    table = s3.shapes.add_table(2, 2, Inches(0.5), Inches(4.5), Inches(5), Inches(1)).table
    for (r, c), text in {(0, 0): "Name", (0, 1): "Wert", (1, 0): "Äpfel", (1, 1): "3"}.items():
        table.cell(r, c).text = text
    prs.save(path)
    return Emu(prs.slide_width), Emu(prs.slide_height)


def make_docx_rich(path):
    """A DOCX with the formatting the direct reader must keep (python-docx plus some raw XML)."""
    import io as _io

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from PIL import Image

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(2)
    section.header.paragraphs[0].text = "Kopfzeile Projekt X"
    section.footer.paragraphs[0].text = "Fußzeile – vertraulich"
    heading_style = doc.styles["Heading 1"]
    heading_style.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    heading_style.font.size = Pt(20)

    doc.add_paragraph("Dokumenttitel", style="Title")
    doc.add_heading("Einleitung", level=1)
    p = doc.add_paragraph()
    for text, props in (("Rot 14pt ", dict(color="FF0000", size=14)), ("Arial ", dict(font="Arial")),
                        ("markiert", dict(highlight=True)), (" E=mc", {}), ("2", dict(superscript=True)),
                        (" durchgestrichen", dict(strike=True))):
        run = p.add_run(text)
        if "color" in props:
            run.font.color.rgb = RGBColor.from_string(props["color"])
        if "size" in props:
            run.font.size = Pt(props["size"])
        if "font" in props:
            run.font.name = props["font"]
        if props.get("highlight"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if props.get("superscript"):
            run.font.superscript = True
        if props.get("strike"):
            run.font.strike = True
    centered = doc.add_paragraph("Zentriert mit Abstand")
    centered.alignment = WD_ALIGN_PARAGRAPH.CENTER
    centered.paragraph_format.space_before = Pt(12)
    centered.paragraph_format.space_after = Pt(18)
    right = doc.add_paragraph("Rechtsbündig")
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    justified = doc.add_paragraph("Blocksatz und eingerückt. " * 8)
    justified.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    justified.paragraph_format.left_indent = Cm(1)
    justified.paragraph_format.line_spacing = 1.5

    doc.add_heading("Listen", level=2)
    doc.add_paragraph("Punkt eins", style="List Bullet")
    doc.add_paragraph("Unterpunkt", style="List Bullet 2")
    doc.add_paragraph("Punkt zwei", style="List Bullet")
    doc.add_paragraph("Schritt eins", style="List Number")
    doc.add_paragraph("Schritt zwei", style="List Number")

    table = doc.add_table(rows=3, cols=3)
    table.style = doc.styles["Table Grid"]
    for r, row in enumerate([("Name", "Menge", "Preis"), ("Äpfel", "3", "1,50"), ("Summe", "", "4,50")]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    for c in range(3):  # header shading
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "D9E2F3")
        table.cell(0, c)._tc.get_or_add_tcPr().append(shd)
    table.cell(2, 0).merge(table.cell(2, 1))  # horizontal merge
    table.cell(2, 0).text = "Summe"

    image = _io.BytesIO()
    Image.new("RGB", (240, 120), (30, 120, 200)).save(image, "PNG")
    image.seek(0)
    doc.add_picture(image, width=Cm(6))

    p = doc.add_paragraph("Mehr unter ")
    rel = p.part.relate_to("https://example.org/info", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "example.org"
    r.append(t)
    link.append(r)
    p._p.append(link)

    doc.add_page_break()
    doc.add_paragraph("Zweite Seite")
    doc.save(path)
    return path


def make_wheel(directory, name="suite_hello", version="1.0"):
    """A minimal pure-Python wheel (py3-none-any) with a module that has a greet() function."""
    path = directory / f"{name}-{version}-py3-none-any.whl"
    dist = f"{name}-{version}.dist-info"
    files = {
        f"{name}/__init__.py": 'def greet(who):\n    return f"Hallo {who} aus dem Wheel"\n',
        f"{dist}/METADATA": f"Metadata-Version: 2.1\nName: {name}\nVersion: {version}\n",
        f"{dist}/WHEEL": "Wheel-Version: 1.0\nGenerator: test\nRoot-Is-Purelib: true\nTag: py3-none-any\n",
    }
    files[f"{dist}/RECORD"] = "".join(f"{n},,\n" for n in files) + f"{dist}/RECORD,,\n"
    with zipfile.ZipFile(path, "w") as z:
        for n, content in files.items():
            z.writestr(n, content)
    return path


def make_pdf(path, pages=3, password=None):
    """A PDF written by hand: pages with Helvetica text (not embedded), an outline, a link from page 1 to the
    last page and a form text field "name" on page 1. With password, the file is encrypted with pypdf."""
    objects = []  # index i -> bytes of object i+1

    def add(body):
        objects.append(body)
        return len(objects)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    pages_id = len(objects) + 1 + 0  # placeholder, fixed below
    page_ids, content_ids = [], []
    reserved_pages = add(b"")  # Pages object, written later
    field = None
    for n in range(1, pages + 1):
        text = f"BT /F1 24 Tf 72 720 Td (Seite {n} - Kapitel {n}) Tj 0 -40 Td /F1 14 Tf (Suchwort Testtext auf Seite {n}) Tj ET"
        if n == 1:
            text += " BT /F1 12 Tf 72 600 Td (Zur letzten Seite) Tj ET"
        content = add(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(text), text.encode("latin-1")))
        content_ids.append(content)
        page_ids.append(add(b""))  # page object, written later
    last = page_ids[-1]
    link = add(b"<< /Type /Annot /Subtype /Link /Rect [72 590 200 615] /Border [0 0 0] /Dest [%d 0 R /XYZ 0 792 0] >>" % last)
    field = add(b"<< /Type /Annot /Subtype /Widget /FT /Tx /T (name) /V () /Rect [72 500 300 525] /F 4 /P %d 0 R /DA (/F1 12 Tf 0 g) /MK << /BC [0 0 0] >> >>" % page_ids[0])
    for i, (pid, cid) in enumerate(zip(page_ids, content_ids)):
        annots = b" /Annots [%d 0 R %d 0 R]" % (link, field) if i == 0 else b""
        objects[pid - 1] = b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R%s >>" % (reserved_pages, font, cid, annots)
    objects[reserved_pages - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % p for p in page_ids), pages)
    outline_root = add(b"")
    items = [add(b"") for _ in page_ids]
    for i, (item, pid) in enumerate(zip(items, page_ids)):
        prev = b" /Prev %d 0 R" % items[i - 1] if i else b""
        nxt = b" /Next %d 0 R" % items[i + 1] if i + 1 < len(items) else b""
        objects[item - 1] = b"<< /Title (Kapitel %d) /Parent %d 0 R /Dest [%d 0 R /Fit]%s%s >>" % (i + 1, outline_root, pid, prev, nxt)
    objects[outline_root - 1] = b"<< /Type /Outlines /First %d 0 R /Last %d 0 R /Count %d >>" % (items[0], items[-1], len(items))
    acroform = add(b"<< /Fields [%d 0 R] /DA (/F1 12 Tf 0 g) /DR << /Font << /F1 %d 0 R >> >> /NeedAppearances true >>" % (field, font))
    catalog = add(b"<< /Type /Catalog /Pages %d 0 R /Outlines %d 0 R /PageMode /UseOutlines /AcroForm %d 0 R >>" % (reserved_pages, outline_root, acroform))
    out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (i, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, catalog, xref)
    path.write_bytes(bytes(out))
    if password:
        from pypdf import PdfReader, PdfWriter
        writer = PdfWriter(clone_from=PdfReader(path))
        writer.encrypt(password, algorithm="AES-256")
        writer.write(path)
    return path
