"""Hardening: opening a malicious file must not run code in the app. Every app gets files whose names, texts, labels, links and
markup carry payloads (tests/helpers/hostile.py); after opening them and looking at the content (selecting, previewing,
presenting), no payload may have run and no executable markup may be in the page or in a frame that can run scripts."""
import base64
import email.message
import io
import json
import sqlite3
import zipfile

import pytest
from playwright.sync_api import expect

from helpers.hostile import HTML, JS, NAME, TEXT, URL, assert_safe, drop_file, open_with_dialog

NO_PICKERS = "window.showOpenFilePicker = undefined; window.showSaveFilePicker = undefined; window.showDirectoryPicker = undefined;"


def zip_bytes(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


def png_bytes(color="red", size=(8, 8)):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


SVG = f'<svg xmlns="http://www.w3.org/2000/svg" width="20" height="20" onload="top.__pwned=1"><script>top.__pwned=1</script><a href="{URL}"><rect width="20" height="20"/></a></svg>'


# ---- text, code and files -----------------------------------------------------------------------------------------------

def test_text_and_hex_editor(app_page):
    for app in ("texteditor", "hexeditor"):
        page = app_page(app, init_script=NO_PICKERS)
        for name, data in ((NAME + ".html", HTML), ("x.md", f"# {TEXT}\n\n{HTML}\n[link]({URL})"), ("x.svg", SVG)):
            drop_file(page, name, data.encode())
        page.wait_for_timeout(500)
        assert_safe(page)


def test_archive_names_and_previews(app_page):
    page = app_page("archive", init_script=NO_PICKERS)
    data = zip_bytes({f"{NAME}/{TEXT}.txt": HTML, "page.html": HTML, "image.svg": SVG, "x.md": f"[a]({URL})"})
    drop_file(page, "evil.zip", data)
    expect(page.locator("[data-testid=row]").first).to_be_visible()
    for name in ("page.html", "image.svg", "x.md"):  # previews
        page.locator("[data-testid=row]", has_text=name).dblclick()
        page.wait_for_timeout(300)
        page.keyboard.press("Escape")
    page.locator("[data-testid=row]").first.dblclick()  # the folder with the odd names
    page.wait_for_timeout(300)
    assert_safe(page)


# ---- office ----------------------------------------------------------------------------------------------------------------

def test_diagram_labels_links_and_images(app_page):
    page = app_page("diagram", init_script=NO_PICKERS)
    label = HTML.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")
    xml = (f'<mxfile><diagram name="{TEXT.replace(chr(34), "&quot;").replace("<", "&lt;").replace(">", "&gt;")}"><mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>'
           f'<mxCell id="a" value="{label}" style="rounded=1;html=1;" vertex="1" parent="1"><mxGeometry x="10" y="10" width="120" height="60" as="geometry"/></mxCell>'
           f'<UserObject id="u" label="{label}" link="{URL}"><mxCell style="text;html=1;" vertex="1" parent="1"><mxGeometry x="200" y="10" width="80" height="30" as="geometry"/></mxCell></UserObject>'
           f'<mxCell id="i" style="shape=image;image={URL};" vertex="1" parent="1"><mxGeometry x="10" y="100" width="40" height="40" as="geometry"/></mxCell>'
           f'<mxCell id="e" value="{label}" edge="1" parent="1" source="a" target="u"><mxGeometry relative="1" as="geometry"/></mxCell>'
           '</root></mxGraphModel></diagram></mxfile>')
    drop_file(page, NAME + ".drawio", xml.encode())
    expect(page.locator("[data-testid=canvas] .dg-shape").first).to_be_visible()
    page.click("[data-testid=sidebar-elements]")
    page.locator("[data-testid=canvas] .dg-shape").first.dblclick(force=True)  # (on a phone the edge label lies over it)
    page.keyboard.press("Escape")
    excalidraw = {"type": "excalidraw", "elements": [{"id": "t", "type": "text", "x": 0, "y": 0, "width": 50, "height": 20, "text": HTML},
                                                     {"id": "i", "type": "image", "x": 0, "y": 40, "width": 20, "height": 20, "fileId": "f"}],
                  "files": {"f": {"dataURL": URL}}}
    page.evaluate("(j) => DG.setDoc(DG.readExcalidraw(j), 'x.excalidraw')", json.dumps(excalidraw))
    page.wait_for_timeout(300)
    assert_safe(page)


def test_document_text_links_and_print(app_page, tmp_path):
    import docx
    d = docx.Document()
    p = d.add_paragraph(TEXT)
    # a hyperlink to javascript:
    part = d.part
    r_id = part.relate_to(URL, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "click me"
    run.append(text)
    link.append(run)
    p._p.append(link)
    path = tmp_path / "evil.docx"
    d.save(path)
    page = app_page("document", init_script=NO_PICKERS + "window.print = () => {};")
    open_with_dialog(page, path)
    expect(page.locator("[data-testid=file-name]")).to_contain_text("evil.docx", timeout=20000)
    page.wait_for_timeout(1500)
    page.click("[data-testid=print]")  # the printed HTML
    page.wait_for_timeout(800)
    assert_safe(page)
    # an HTML file opened as a document
    (tmp_path / "evil.html").write_text(f"<html><body><p>{HTML}</p><a href='{URL}'>x</a></body></html>")
    page.evaluate("DOC.state.dirty = false")
    open_with_dialog(page, tmp_path / "evil.html")
    page.wait_for_timeout(2000)
    page.click("[data-testid=print]")
    page.wait_for_timeout(800)
    assert_safe(page)


def test_spreadsheet_cells_names_and_links(app_page, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "x<img onerror"
    ws["A1"] = TEXT
    ws["A2"] = HTML
    ws["A3"] = "link"
    ws["A3"].hyperlink = URL
    path = tmp_path / "evil.xlsx"
    wb.save(path)
    page = app_page("spreadsheet", init_script=NO_PICKERS)
    open_with_dialog(page, path)
    expect(page.locator("[data-testid=file-name]")).to_contain_text("evil.xlsx", timeout=20000)
    page.wait_for_timeout(1500)
    assert_safe(page)
    (tmp_path / "evil.csv").write_text(f'"{TEXT}",=HYPERLINK("{URL}")\n')
    open_with_dialog(page, tmp_path / "evil.csv")
    page.wait_for_timeout(500)
    if page.locator("[data-testid=csv-ok]").count():
        page.click("[data-testid=csv-ok]")
    page.wait_for_timeout(1000)
    assert_safe(page)


def test_presentation_texts_links_and_console(app_page, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = TEXT
    run = slide.placeholders[1].text_frame.paragraphs[0].add_run()
    run.text = HTML
    run.hyperlink.address = URL
    slide.notes_slide.notes_text_frame.text = HTML
    box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1))
    box.name = NAME
    box.text_frame.text = TEXT
    path = tmp_path / "evil.pptx"
    prs.save(path)
    page = app_page("presentation", init_script=NO_PICKERS)
    drop_file(page, "evil.pptx", path.read_bytes())
    page.wait_for_function("() => PR.state.name === 'evil.pptx'", timeout=20000)
    page.locator("[data-testid=panel]").first.wait_for(state="attached")  # (a drawer on phones)
    assert_safe(page)


def test_web_slides_html_markdown_and_presenting(app_page, tmp_path):
    page = app_page("reveal", init_script=NO_PICKERS)
    deck = (f"<!doctype html><html><head><title>{TEXT}</title></head><body><div class='reveal'><div class='slides'>"
            f"<section><h2>{TEXT}</h2>{HTML}<a href='{URL}'>x</a></section>"
            f"<section data-markdown><textarea data-template># {TEXT}\n\n{HTML}\n\n[x]({URL})</textarea></section>"
            f"<section data-background-image='{URL}' data-background-iframe='data:text/html,&lt;script&gt;top.__pwned=1&lt;/script&gt;'><p onclick='top.__pwned=1'>p</p><aside class='notes'>{HTML}</aside></section>"
            "</div></div></body></html>")
    drop_file(page, "evil.html", deck.encode())
    # the user is asked whether the file's scripts may run; by default they stay blocked
    expect(page.locator("[data-testid=confirm]")).to_contain_text("javascript:")
    page.click("[data-testid=confirm-cancel]")
    expect(page.locator("[data-testid=scripts]")).to_be_visible()
    page.wait_for_timeout(2500)
    assert_safe(page)
    md = f"# {TEXT}\n\n{HTML}\n\n[x]({URL})\n\n---\n\n<script>top.__pwned=1</script>\n"
    drop_file(page, "evil.md", md.encode())
    page.click("[data-testid=confirm-cancel]")
    page.wait_for_timeout(2500)
    assert_safe(page)
    # presenting (in the page and in a new window) and the saved file keep them blocked as well
    page.click("[data-testid=present]")
    page.wait_for_timeout(1500)
    assert_safe(page)
    page.keyboard.press("Escape")
    page.evaluate("document.fullscreenElement && document.exitFullscreen()")
    with page.context.expect_page() as popup:
        page.evaluate("document.querySelector('[data-testid=present-window]').click()")
    win = popup.value
    win.wait_for_timeout(1500)
    assert_safe(page, win)
    saved = page.evaluate("RV.presentationHtml()")
    assert "Content-Security-Policy" in saved


def test_web_slides_trusted_scripts(app_page):
    """A presentation with its own scripts: they run only after the user allowed them, and can be blocked again."""
    page = app_page("reveal", init_script=NO_PICKERS)
    deck = ("<!doctype html><html><body><div class='reveal'><div class='slides'>"
            "<section><h2 id='t'>x</h2><script>top.__pwned=(top.__pwned||0)+1</script></section></div></div></body></html>")
    drop_file(page, "own.html", deck.encode())
    dialog = page.locator("[data-testid=confirm]")
    expect(dialog).to_contain_text("unprotected")
    page.click("[data-testid=confirm-ok]")
    expect(page.locator("[data-testid=scripts]")).to_have_class("rv-trusted")
    page.wait_for_function("() => window.__pwned > 0", timeout=10000)
    assert "Content-Security-Policy" not in page.evaluate("RV.presentationHtml()")
    # blocked again
    page.click("[data-testid=scripts]")
    page.evaluate("window.__pwned = 0")
    page.wait_for_timeout(1500)
    assert page.evaluate("window.__pwned || 0") == 0
    assert "Content-Security-Policy" in page.evaluate("RV.presentationHtml()")
    # asked again before allowing them once more
    page.click("[data-testid=scripts]")
    expect(dialog).to_be_visible()
    page.click("[data-testid=confirm-cancel]")
    assert [e for e in page.errors if "Content Security Policy" not in e and "Content-Security-Policy" not in e] == []  # (the blocked script is reported)


def test_notebook_outputs_and_markdown(app_page):
    nb = {"nbformat": 4, "nbformat_minor": 5, "metadata": {"kernelspec": {"name": "python3", "display_name": "Python 3"}},
          "cells": [
              {"cell_type": "markdown", "id": "m", "metadata": {}, "source": [f"# {TEXT}\n", HTML, f"\n[x]({URL})"]},
              {"cell_type": "code", "id": "c", "metadata": {}, "execution_count": 1, "source": ["1"], "outputs": [
                  {"output_type": "display_data", "metadata": {}, "data": {"text/html": [HTML], "image/svg+xml": [SVG], "text/markdown": [HTML],
                                                                           "application/javascript": ["top.__pwned=1"], "text/plain": ["x"]}},
                  {"output_type": "execute_result", "execution_count": 1, "metadata": {}, "data": {"text/html": [f"<table><tr><td>{HTML}</td></tr></table>"]}},
                  {"output_type": "error", "ename": TEXT, "evalue": HTML, "traceback": [HTML]},
                  {"output_type": "stream", "name": "stdout", "text": [HTML]}]}]}
    page = app_page("notebook", init_script=NO_PICKERS)
    drop_file(page, NAME + ".ipynb", json.dumps(nb).encode())
    expect(page.locator("[data-testid=output]").first).to_be_visible(timeout=20000)
    assert_safe(page)


def test_pdf_javascript_links_and_outline(app_page, tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, TextStringObject, ArrayObject, FloatObject, NumberObject
    w = PdfWriter()
    w.add_blank_page(300, 300)
    w.add_js("app.alert(1); this.__pwned = 1;")
    w.add_metadata({"/Title": TEXT})
    w.add_outline_item(TEXT, 0)
    annot = DictionaryObject({NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Link"),
                              NameObject("/Rect"): ArrayObject([FloatObject(10), FloatObject(10), FloatObject(200), FloatObject(100)]),
                              NameObject("/Border"): ArrayObject([NumberObject(0), NumberObject(0), NumberObject(0)]),
                              NameObject("/A"): DictionaryObject({NameObject("/S"): NameObject("/URI"), NameObject("/URI"): TextStringObject(URL)})})
    w.pages[0][NameObject("/Annots")] = ArrayObject([w._add_object(annot)])
    buf = io.BytesIO()
    w.write(buf)
    page = app_page("pdfviewer", init_script=NO_PICKERS)
    drop_file(page, NAME + ".pdf", buf.getvalue())
    page.wait_for_timeout(2500)
    assert_safe(page)


# ---- organiser and security ------------------------------------------------------------------------------------------------

def test_calendar_event_texts_and_urls(app_page):
    ics = ("BEGIN:VCALENDAR\r\nVERSION:2.0\r\nPRODID:x\r\nX-WR-CALNAME:" + TEXT + "\r\nBEGIN:VEVENT\r\nUID:a@x\r\nDTSTART:20261014T100000Z\r\n"
           "DTEND:20261014T110000Z\r\nSUMMARY:" + TEXT + "\r\nDESCRIPTION:" + HTML.replace(",", "\\,").replace(";", "\\;") + "\r\nLOCATION:" + TEXT +
           "\r\nURL:" + URL + "\r\nATTACH:" + URL + "\r\nORGANIZER;CN=" + "\"x<img onerror\"" + ":mailto:" + URL + "\r\nEND:VEVENT\r\nEND:VCALENDAR\r\n")
    page = app_page("calendar", init_script=NO_PICKERS + "Date.now = () => new Date('2026-10-14T12:00:00Z').getTime();")
    drop_file(page, NAME + ".ics", ics.encode())
    page.wait_for_timeout(1000)
    if page.locator("[data-testid=event]").count():
        page.locator("[data-testid=event]").first.dblclick()
        page.wait_for_timeout(300)
    assert_safe(page)


def test_mail_messages_and_contacts(app_page):
    msg = email.message.EmailMessage()
    msg["Subject"] = TEXT.replace('"', "'")
    msg["From"] = "<img src=x onerror=top.__pwned=1> <a@example.com>"
    msg["To"] = "b@example.com"
    msg.set_content("plain " + TEXT)
    msg.add_alternative(f"<html><head><meta http-equiv='refresh' content='0;url={URL}'><base href='{URL}'></head><body>{HTML}"
                        f"<a href='{URL}'>x</a><form action='{URL}'><button>b</button></form><object data='{URL}'></object></body></html>", subtype="html")
    msg.add_attachment(HTML.encode(), maintype="text", subtype="html", filename=NAME + ".html")
    msg.add_attachment(SVG.encode(), maintype="image", subtype="svg+xml", filename="x.svg")
    page = app_page("mail", init_script=NO_PICKERS)
    drop_file(page, "evil.eml", bytes(msg))
    expect(page.locator("[data-testid=message]").first).to_be_visible(timeout=15000)
    page.locator("[data-testid=message]").first.click()
    page.wait_for_timeout(1000)
    assert_safe(page)
    vcf = f"BEGIN:VCARD\r\nVERSION:3.0\r\nFN:{TEXT}\r\nEMAIL:x@example.com\r\nURL:{URL}\r\nNOTE:{HTML}\r\nEND:VCARD\r\n"
    drop_file(page, "evil.vcf", vcf.encode())
    page.wait_for_timeout(800)
    assert_safe(page)


def test_keepass_entries(app_page, tmp_path):
    from pykeepass import create_database
    path = tmp_path / "evil.kdbx"
    kp = create_database(str(path), password="pw")
    odd = "<img src=x onerror=top.__pwned=1>"  # (pykeepass cannot handle quotes in names)
    group = kp.add_group(kp.root_group, odd)
    entry = kp.add_entry(group, odd, odd, "pw", url=URL, notes=HTML)
    entry.set_custom_property("x" + odd, HTML)
    kp.save()
    page = app_page("keepass", init_script=NO_PICKERS)
    drop_file(page, "evil.kdbx", path.read_bytes())
    page.fill("[data-testid=password]", "pw")
    page.keyboard.press("Enter")
    expect(page.locator("[data-testid=group]").first).to_be_visible(timeout=20000)
    page.locator("[data-testid=group]", has_text="img").first.click()
    page.locator("[data-testid=entry]").first.click()
    page.wait_for_timeout(500)
    assert_safe(page)


def test_database_names_and_values(app_page, tmp_path):
    path = tmp_path / "evil.sqlite"
    con = sqlite3.connect(path)
    q = '"' + TEXT.replace('"', '""') + '"'  # a quoted identifier
    con.execute(f'CREATE TABLE {q} ({q} TEXT, "x" TEXT)')
    con.execute(f'INSERT INTO {q} VALUES (?, ?)', (HTML, URL))
    con.commit()
    con.close()
    page = app_page("dbtool", init_script=NO_PICKERS)
    drop_file(page, "evil.sqlite", path.read_bytes())
    expect(page.locator("[data-testid=table]").first).to_be_visible(timeout=20000)
    page.locator("[data-testid=table]").first.locator("summary").click()
    page.locator("[data-testid=browse]").first.click()
    page.wait_for_timeout(800)
    assert_safe(page)


# ---- media and web ---------------------------------------------------------------------------------------------------------

def test_images(app_page):
    page = app_page("imageeditor", init_script=NO_PICKERS)
    drop_file(page, NAME + ".svg", SVG.encode(), "image/svg+xml")
    page.wait_for_timeout(1000)
    ora = zip_bytes({"mimetype": "image/openraster", "stack.xml": f'<image w="8" h="8"><stack><layer name="{TEXT.replace(chr(34), "&quot;").replace("<", "&lt;").replace(">", "&gt;")}" src="data/a.png"/></stack></image>',
                     "data/a.png": png_bytes()})
    drop_file(page, "evil.ora", ora)
    page.wait_for_timeout(1000)
    assert_safe(page)
    page = app_page("barcode", init_script=NO_PICKERS)
    drop_file(page, NAME + ".svg", SVG.encode(), "image/svg+xml")
    page.wait_for_timeout(800)
    assert_safe(page)


def test_media_player_playlists(app_page):
    from helpers.media_samples import make_wav
    page = app_page("mediaplayer", init_script=NO_PICKERS)
    playlist = f"#EXTM3U\n#EXTINF:5,{TEXT}\n{URL}\n#EXTINF:5,{HTML}\nhttps://example.com/{NAME}.mp3\n"
    drop_file(page, NAME + ".m3u", playlist.encode(), "audio/x-mpegurl")
    page.wait_for_timeout(800)
    assert_safe(page)


def test_web_inspector_pages_and_har(app_page, tmp_path):
    page = app_page("webinspector", init_script=NO_PICKERS)
    doc = f"<!doctype html><html><head><title>{TEXT}</title><meta http-equiv='refresh' content='0;url={URL}'><script>top.__pwned=1</script></head><body>{HTML}<a href='{URL}'>x</a></body></html>"
    drop_file(page, NAME + ".html", doc.encode(), "text/html")
    page.wait_for_timeout(1500)
    for i in range(min(6, page.locator("[role=tab]").count())):
        tabs = page.locator("[role=tab]")
        if i < tabs.count():
            tabs.nth(i).click()
            page.wait_for_timeout(150)
    assert_safe(page)
    har = {"log": {"version": "1.2", "creator": {"name": TEXT, "version": "1"}, "entries": [{
        "startedDateTime": "2026-10-14T10:00:00Z", "time": 5, "request": {"method": TEXT, "url": URL, "httpVersion": "HTTP/1.1", "headers": [{"name": TEXT, "value": HTML}], "cookies": [], "queryString": [], "headersSize": -1, "bodySize": 0},
        "response": {"status": 200, "statusText": TEXT, "httpVersion": "HTTP/1.1", "headers": [{"name": "content-type", "value": "text/html"}], "cookies": [],
                     "content": {"size": 10, "mimeType": "text/html", "text": HTML}, "redirectURL": URL, "headersSize": -1, "bodySize": 10},
        "cache": {}, "timings": {"send": 0, "wait": 5, "receive": 0}}]}}
    drop_file(page, NAME + ".har", json.dumps(har).encode(), "application/json")
    page.wait_for_timeout(1000)
    for row in page.locator("[data-testid=entry], [data-testid=request]").all()[:2]:
        row.click()
        page.wait_for_timeout(200)
    assert_safe(page)


def test_http_collections(app_page, tmp_path):
    collection = {"info": {"name": TEXT, "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json"},
                  "item": [{"name": TEXT, "request": {"method": "GET", "url": URL, "header": [{"key": TEXT, "value": HTML}], "description": HTML},
                            "event": [{"listen": "test", "script": {"exec": ["top.__pwned=1; self.__pwned=1;"]}}]}]}
    path = tmp_path / "evil.postman_collection.json"
    path.write_text(json.dumps(collection))
    page = app_page("http", init_script=NO_PICKERS)
    open_with_dialog(page, path)
    expect(page.locator("[data-testid=tree-request]")).to_have_count(1)
    page.locator("[data-testid=tree-request]").first.locator("a").click()
    page.wait_for_timeout(500)
    assert_safe(page)


@pytest.mark.parametrize("app", ["document", "spreadsheet", "presentation", "pdfviewer", "mail", "webinspector"])
def test_script_links_are_refused(app_page, app):
    """Whatever component opens a link from a file (Univer's link popups, viewers): javascript: never runs in the app."""
    page = app_page(app, init_script=NO_PICKERS)
    assert page.evaluate(f"window.open({URL!r}) === null && window.open(' JaVaScRiPt:{JS}', '_blank') === null")
    page.evaluate(f"""() => {{ const a = document.createElement('a'); a.href = {URL!r}; a.textContent = 'x'; a.id = 'evil';
      document.body.append(a); a.click(); a.dispatchEvent(new MouseEvent('click', {{ bubbles: true, cancelable: true }})); }}""")
    page.wait_for_timeout(300)
    assert page.evaluate("window.__pwned || 0") == 0
